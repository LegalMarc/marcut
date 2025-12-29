
from io import BytesIO
import copy
import os
import posixpath
import re
import zipfile
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.text.run import Run
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from datetime import datetime, timezone
from typing import List, Dict, Any, Optional, Iterable, Tuple
from dataclasses import dataclass, field, fields

def _safe_fromstring(xml_bytes: bytes):
    """Safe XML parsing that disables entity resolution."""
    from lxml import etree
    parser = etree.XMLParser(resolve_entities=False)
    return etree.fromstring(xml_bytes, parser)


CLI_ARG_PAIRS: List[Tuple[str, str]] = [
    # App Properties
    ("--no-clean-company", "clean_company"),
    ("--no-clean-manager", "clean_manager"),
    ("--no-clean-editing-time", "clean_total_editing_time"),
    ("--no-clean-application", "clean_application"),
    ("--no-clean-app-version", "clean_app_version"),
    ("--no-clean-template", "clean_template"),
    ("--no-clean-hyperlink-base", "clean_hyperlink_base"),
    ("--no-clean-statistics", "clean_statistics"),
    ("--no-clean-doc-security", "clean_doc_security"),
    ("--no-clean-scale-crop", "clean_scale_crop"),
    ("--no-clean-links-up-to-date", "clean_links_up_to_date"),
    ("--no-clean-shared-doc", "clean_shared_doc"),
    ("--no-clean-hyperlinks-changed", "clean_hyperlinks_changed"),
    # Core Properties
    ("--no-clean-author", "clean_author"),
    ("--no-clean-last-modified-by", "clean_last_modified_by"),
    ("--no-clean-title", "clean_title"),
    ("--no-clean-subject", "clean_subject"),
    ("--no-clean-keywords", "clean_keywords"),
    ("--no-clean-comments", "clean_comments"),
    ("--no-clean-category", "clean_category"),
    ("--no-clean-content-status", "clean_content_status"),
    ("--no-clean-created-date", "clean_created_date"),
    ("--no-clean-modified-date", "clean_modified_date"),
    ("--no-clean-last-printed", "clean_last_printed"),
    ("--no-clean-revision", "clean_revision_number"),
    ("--no-clean-identifier", "clean_identifier"),
    ("--no-clean-language", "clean_language"),
    ("--no-clean-version", "clean_version"),
    # Custom Properties
    ("--no-clean-custom-props", "clean_custom_properties"),
    # Document Structure
    ("--no-clean-review-comments", "clean_review_comments"),
    ("--no-clean-track-changes", "clean_track_changes"),
    ("--no-clean-rsids", "clean_rsids"),
    ("--no-clean-guid", "clean_document_guid"),
    ("--no-clean-spell-grammar", "clean_spell_grammar_state"),
    ("--no-clean-doc-vars", "clean_document_variables"),
    ("--no-clean-mail-merge", "clean_mail_merge"),
    ("--no-clean-data-bindings", "clean_data_bindings"),
    ("--no-clean-doc-versions", "clean_document_versions"),
    ("--no-clean-ink-annotations", "clean_ink_annotations"),
    ("--no-clean-hidden-text", "clean_hidden_text"),
    ("--no-clean-invisible-objects", "clean_invisible_objects"),
    ("--no-clean-headers-footers", "clean_headers_footers"),
    ("--no-clean-watermarks", "clean_watermarks"),
    # Embedded Content
    ("--no-clean-thumbnail", "clean_thumbnail"),
    ("--no-clean-hyperlinks", "clean_hyperlink_urls"),
    ("--no-clean-alt-text", "clean_alt_text"),
    ("--no-clean-ole", "clean_ole_objects"),
    ("--no-clean-macros", "clean_vba_macros"),
    ("--no-clean-signatures", "clean_digital_signatures"),
    ("--no-clean-printer", "clean_printer_settings"),
    ("--no-clean-fonts", "clean_embedded_fonts"),
    ("--no-clean-glossary", "clean_glossary"),
    ("--no-clean-fast-save", "clean_fast_save_data"),
    # Advanced Hardening
    ("--no-clean-ext-links", "clean_external_links"),
    ("--no-clean-unc-paths", "clean_unc_paths"),
    ("--no-clean-user-paths", "clean_user_paths"),
    ("--no-clean-internal-urls", "clean_internal_urls"),
    ("--no-clean-ole-sources", "clean_ole_sources"),
    ("--no-clean-exif", "clean_image_exif"),
    ("--no-clean-style-names", "clean_style_names"),
    ("--no-clean-chart-labels", "clean_chart_labels"),
    ("--no-clean-form-defaults", "clean_form_defaults"),
    ("--no-clean-language-settings", "clean_language_settings"),
    ("--no-clean-activex", "clean_activex"),
]

CLI_ARG_MAP = dict(CLI_ARG_PAIRS)
FIELD_TO_CLI = {field: flag for flag, field in CLI_ARG_PAIRS}


@dataclass
class MetadataCleaningSettings:
    """Settings for granular control of which metadata fields are cleaned during redaction."""
    
    # App Properties (docProps/app.xml)
    clean_company: bool = True
    clean_manager: bool = True
    clean_total_editing_time: bool = True
    clean_application: bool = True
    clean_app_version: bool = True
    clean_template: bool = True
    clean_hyperlink_base: bool = True
    clean_statistics: bool = True  # chars, words, lines, paragraphs, pages
    clean_doc_security: bool = True
    clean_scale_crop: bool = True
    clean_links_up_to_date: bool = True
    clean_shared_doc: bool = True
    clean_hyperlinks_changed: bool = True
    
    # Core Properties (docProps/core.xml)
    clean_author: bool = True
    clean_last_modified_by: bool = True
    clean_title: bool = True
    clean_subject: bool = True
    clean_keywords: bool = True
    clean_comments: bool = True
    clean_category: bool = True
    clean_content_status: bool = True
    clean_created_date: bool = False  # Default OFF per user request
    clean_modified_date: bool = False  # Default OFF per user request
    clean_last_printed: bool = True
    clean_revision_number: bool = True
    clean_identifier: bool = True
    clean_language: bool = True
    clean_version: bool = True
    
    # Custom Properties
    clean_custom_properties: bool = True
    
    # Document Structure
    clean_review_comments: bool = True
    clean_track_changes: bool = True
    clean_rsids: bool = True
    clean_document_guid: bool = True
    clean_spell_grammar_state: bool = True
    clean_document_variables: bool = True
    clean_mail_merge: bool = True
    clean_data_bindings: bool = True
    clean_document_versions: bool = True
    clean_ink_annotations: bool = True
    clean_hidden_text: bool = True
    clean_invisible_objects: bool = True
    clean_headers_footers: bool = True
    clean_watermarks: bool = True
    
    # Embedded Content
    clean_thumbnail: bool = True
    clean_hyperlink_urls: bool = True
    clean_alt_text: bool = True
    clean_ole_objects: bool = True
    clean_vba_macros: bool = True
    clean_digital_signatures: bool = True
    clean_printer_settings: bool = True
    clean_embedded_fonts: bool = True
    clean_glossary: bool = True
    clean_fast_save_data: bool = True

    # Advanced Hardening
    clean_external_links: bool = True
    clean_unc_paths: bool = True
    clean_user_paths: bool = True
    clean_internal_urls: bool = True
    clean_ole_sources: bool = True
    clean_image_exif: bool = True
    clean_style_names: bool = True
    clean_chart_labels: bool = True
    clean_form_defaults: bool = True
    clean_language_settings: bool = True
    clean_activex: bool = True
    
    @classmethod
    def from_cli_args(cls, args: List[str]) -> "MetadataCleaningSettings":
        """Parse CLI arguments to create settings. Args are in format --no-clean-XYZ."""
        settings = cls()
        for arg in args:
            if arg in CLI_ARG_MAP:
                setattr(settings, CLI_ARG_MAP[arg], False)
        return settings

    def to_cli_args(self) -> List[str]:
        """Generate CLI arguments that disable any settings set to False."""
        args: List[str] = []
        if all(not getattr(self, f.name) for f in fields(self)):
            args.append("--preset-none")
        for field_name, flag in FIELD_TO_CLI.items():
            if hasattr(self, field_name) and not getattr(self, field_name):
                args.append(flag)
        return args


class DocxMap:
    def __init__(self, doc: Document, author_name: str = "Marcut"):
        self.doc = doc
        self.text = ""
        self.index = []
        self._rev_id = 1
        self.author_name = author_name
        self.detached_parts = [] # Track pars that need manual blob update
        self._build()

    @staticmethod
    def load(path: str) -> "DocxMap":
        return DocxMap(Document(path))

    @staticmethod
    def load_accepting_revisions(path: str, debug: bool = False) -> "DocxMap":
        from .docx_revisions import accept_revisions_in_docx_bytes

        docx_bytes, changed = accept_revisions_in_docx_bytes(path, debug=debug)
        if not changed or docx_bytes is None:
            return DocxMap.load(path)
        return DocxMap(Document(BytesIO(docx_bytes)))

    def save(self, path: str):
        # Flush detached parts back to blob before saving
        from lxml import etree
        for part, root in self.detached_parts:
            # Manually update the blob because python-docx treats this as a generic Part
            # Ensure proper XML declaration
            part._blob = etree.tostring(root, encoding='UTF-8', xml_declaration=True)
            
        self.doc.save(path)
        self._postprocess_zip(path)

    def _postprocess_zip(self, path: str):
        settings = getattr(self, "_metadata_settings", None)
        if settings is None:
            return
        try:
            self._rewrite_docx_zip(path, settings)
        except Exception:
            # Best-effort hardening; don't fail save on post-processing errors.
            pass

    def _rewrite_docx_zip(self, path: str, settings: MetadataCleaningSettings):
        from lxml import etree

        def _rels_source_dir(rels_path: str) -> str:
            if rels_path == "_rels/.rels":
                return ""
            if "/_rels/" not in rels_path:
                return posixpath.dirname(rels_path)
            source_path = rels_path.replace("/_rels/", "/")
            if source_path.endswith(".rels"):
                source_path = source_path[:-5]
            return posixpath.dirname(source_path)

        def _resolve_target(rels_path: str, target: str) -> str:
            base_dir = _rels_source_dir(rels_path)
            return posixpath.normpath(posixpath.join(base_dir, target))

        def _is_unc_path(target: str) -> bool:
            return target.startswith("\\\\") or (target.startswith("//") and not target.startswith("http"))

        def _is_user_path(target: str) -> bool:
            return ("/Users/" in target or "/home/" in target or "C:\\\\Users\\" in target
                    or "C:/Users/" in target or "%USERPROFILE%" in target)

        def _is_file_path(target: str) -> bool:
            if target.startswith("file:"):
                return True
            if re.match(r"^[A-Za-z]:[\\\\/]", target):
                return True
            return target.startswith("/") or target.startswith("./") or target.startswith("../")

        def _is_internal_url(target: str) -> bool:
            if not target.startswith(("http://", "https://")):
                return False
            try:
                host = target.split("//", 1)[1].split("/", 1)[0]
            except Exception:
                return False
            if host.startswith(("127.", "10.", "192.168.", "169.254.")):
                return True
            if host.endswith(".local") or host.endswith(".lan"):
                return True
            return "." not in host

        def _sanitize_target(target: str, r_type: str) -> str:
            if settings.clean_ole_sources and "oleObject" in r_type:
                return "urn:marcut:redacted"
            if settings.clean_unc_paths and _is_unc_path(target):
                return "urn:marcut:redacted"
            if settings.clean_user_paths and _is_user_path(target):
                return "urn:marcut:redacted"
            if settings.clean_internal_urls and _is_internal_url(target):
                return "urn:marcut:redacted"
            if settings.clean_external_links and _is_file_path(target):
                return "urn:marcut:redacted"
            return target

        def _strip_jpeg_metadata(data: bytes) -> Tuple[bytes, bool]:
            if not data.startswith(b"\xFF\xD8"):
                return data, False
            out = bytearray(b"\xFF\xD8")
            i = 2
            changed = False
            while i + 1 < len(data):
                if data[i] != 0xFF:
                    out.extend(data[i:])
                    break
                # Skip fill bytes
                while i < len(data) and data[i] == 0xFF:
                    i += 1
                if i >= len(data):
                    break
                marker = data[i]
                i += 1
                if marker == 0xDA:  # SOS
                    out.extend(b"\xFF\xDA")
                    out.extend(data[i:])
                    break
                if i + 1 >= len(data):
                    break
                length = (data[i] << 8) + data[i + 1]
                segment_start = i - 1
                segment_end = i + length
                if marker in (0xE1, 0xED):  # APP1/APP13
                    changed = True
                else:
                    out.extend(b"\xFF" + bytes([marker]))
                    out.extend(data[i:segment_end])
                i = segment_end
            return bytes(out), changed

        def _strip_png_metadata(data: bytes) -> Tuple[bytes, bool]:
            if not data.startswith(b"\x89PNG\r\n\x1a\n"):
                return data, False
            out = bytearray(data[:8])
            i = 8
            changed = False
            while i + 8 <= len(data):
                length = int.from_bytes(data[i:i + 4], "big")
                chunk_type = data[i + 4:i + 8]
                chunk_end = i + 12 + length
                chunk = data[i:chunk_end]
                if chunk_type in (b"tEXt", b"iTXt", b"zTXt", b"eXIf"):
                    changed = True
                else:
                    out.extend(chunk)
                i = chunk_end
            return bytes(out), changed

        def _remove_lang_elements(root: etree._Element) -> bool:
            changed = False
            for el in list(root.iter()):
                if el.tag == qn("w:lang"):
                    parent = el.getparent()
                    if parent is not None:
                        parent.remove(el)
                        changed = True
            return changed

        def _strip_form_defaults(root: etree._Element) -> bool:
            changed = False
            for tag in (qn("w:default"), qn("w:result")):
                for el in list(root.iter(tag)):
                    parent = el.getparent()
                    if parent is not None:
                        parent.remove(el)
                        changed = True
            return changed

        def _rename_custom_styles(root: etree._Element) -> bool:
            changed = False
            idx = 1
            for style in root.findall(".//w:style", namespaces=root.nsmap):
                if style.get(qn("w:customStyle")) != "1":
                    continue
                name_el = style.find("w:name", namespaces=root.nsmap)
                if name_el is not None:
                    name_el.set(qn("w:val"), f"CustomStyle{idx}")
                    changed = True
                    idx += 1
                aliases = style.find("w:aliases", namespaces=root.nsmap)
                if aliases is not None:
                    style.remove(aliases)
                    changed = True
            return changed

        def _clean_chart_labels(root: etree._Element) -> bool:
            changed = False
            for el in root.iter():
                if el.tag.endswith("}v") or el.tag.endswith("}t"):
                    if el.text:
                        el.text = "[REDACTED]"
                        changed = True
            return changed

        def _scrub_rels(xml_bytes: bytes, rels_path: str) -> Tuple[bytes, bool]:
            root = _safe_fromstring(xml_bytes)
            changed = False
            if rels_path == "_rels/.rels":
                core_type = "http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties"
                legacy_core_type = "http://schemas.openxmlformats.org/officedocument/2006/relationships/metadata/core-properties"
                core_rels = [
                    rel for rel in list(root)
                    if rel.get("Type") in (core_type, legacy_core_type)
                ]
                if core_rels:
                    has_correct = any(rel.get("Type") == core_type for rel in core_rels)
                    if has_correct:
                        for rel in core_rels:
                            if rel.get("Type") == legacy_core_type:
                                root.remove(rel)
                                changed = True
                    else:
                        primary = core_rels[0]
                        primary.set("Type", core_type)
                        changed = True
                        for rel in core_rels[1:]:
                            root.remove(rel)
                            changed = True
            for rel in list(root):
                r_type = rel.get("Type") or ""
                target = rel.get("Target") or ""
                target_mode = rel.get("TargetMode") or ""
                resolved = _resolve_target(rels_path, target)

                if resolved in removed_parts:
                    root.remove(rel)
                    changed = True
                    continue
                if settings.clean_custom_properties and "customXml" in r_type:
                    root.remove(rel)
                    changed = True
                    continue
                if settings.clean_digital_signatures and "signature" in r_type:
                    root.remove(rel)
                    changed = True
                    continue
                if settings.clean_printer_settings and "printerSettings" in r_type:
                    root.remove(rel)
                    changed = True
                    continue
                if settings.clean_document_versions and "versions" in r_type:
                    root.remove(rel)
                    changed = True
                    continue
                if settings.clean_ink_annotations and "ink" in r_type:
                    root.remove(rel)
                    changed = True
                    continue
                if settings.clean_headers_footers and ("header" in r_type or "footer" in r_type):
                    root.remove(rel)
                    changed = True
                    continue
                if settings.clean_activex and ("control" in r_type or "activeX" in r_type):
                    root.remove(rel)
                    changed = True
                    continue

                if target_mode == "External":
                    sanitized = _sanitize_target(target, r_type)
                    if sanitized != target:
                        rel.set("Target", sanitized)
                        changed = True
            if changed:
                return etree.tostring(root, encoding="UTF-8", xml_declaration=True), True
            return xml_bytes, False

        def _update_content_types(xml_bytes: bytes, removed: Iterable[str]) -> Tuple[bytes, bool]:
            root = _safe_fromstring(xml_bytes)
            removed_parts = {f"/{name}" for name in removed}
            changed = False
            for override in list(root.findall("Override")):
                if override.get("PartName") in removed_parts:
                    root.remove(override)
                    changed = True
            if changed:
                return etree.tostring(root, encoding="UTF-8", xml_declaration=True), True
            return xml_bytes, False

        with zipfile.ZipFile(path, "r") as zin:
            items = zin.infolist()
            last_index = {item.filename: idx for idx, item in enumerate(items)}
            removed_parts = set()
            for item in items:
                name = item.filename
                if settings.clean_thumbnail and name.startswith("docProps/thumbnail"):
                    removed_parts.add(name)
                if settings.clean_custom_properties and (name.startswith("customXml/") or name == "docProps/custom.xml"):
                    removed_parts.add(name)
                if settings.clean_review_comments and name.startswith("word/comments"):
                    removed_parts.add(name)
                if settings.clean_ole_objects and name.startswith("word/embeddings/"):
                    removed_parts.add(name)
                if settings.clean_activex and (name.startswith("word/activeX/") or name.startswith("word/controls/")):
                    removed_parts.add(name)
                if settings.clean_embedded_fonts and name.startswith("word/fonts/"):
                    removed_parts.add(name)
                if settings.clean_headers_footers and (name.startswith("word/header") or name.startswith("word/footer")):
                    removed_parts.add(name)
                if settings.clean_document_versions and name.startswith("word/versions"):
                    removed_parts.add(name)
                if settings.clean_ink_annotations and name.startswith("word/ink"):
                    removed_parts.add(name)

            updates: Dict[str, bytes] = {}

            for item in items:
                name = item.filename
                if name in removed_parts:
                    continue

                data = zin.read(name)

                if name.endswith(".rels"):
                    updated, changed = _scrub_rels(data, name)
                    if changed:
                        updates[name] = updated
                    continue

                if name == "[Content_Types].xml":
                    updated, changed = _update_content_types(data, removed_parts)
                    if changed:
                        updates[name] = updated
                    continue

                if settings.clean_image_exif and name.startswith("word/media/"):
                    lower = name.lower()
                    if lower.endswith((".jpg", ".jpeg")):
                        updated, changed = _strip_jpeg_metadata(data)
                        if changed:
                            updates[name] = updated
                        continue
                    if lower.endswith(".png"):
                        updated, changed = _strip_png_metadata(data)
                        if changed:
                            updates[name] = updated
                        continue

                if settings.clean_embedded_fonts and name == "word/fontTable.xml":
                    root = _safe_fromstring(data)
                    changed = False
                    for tag in ("embedRegular", "embedBold", "embedItalic", "embedBoldItalic"):
                        for el in list(root.findall(f".//w:{tag}", namespaces=root.nsmap)):
                            parent = el.getparent()
                            if parent is not None:
                                parent.remove(el)
                                changed = True
                    if changed:
                        updates[name] = etree.tostring(root, encoding="UTF-8", xml_declaration=True)
                    continue

                if name == "word/styles.xml" and (settings.clean_style_names or settings.clean_language_settings):
                    root = _safe_fromstring(data)
                    changed = False
                    if settings.clean_style_names:
                        changed = _rename_custom_styles(root) or changed
                    if settings.clean_language_settings:
                        changed = _remove_lang_elements(root) or changed
                    if changed:
                        updates[name] = etree.tostring(root, encoding="UTF-8", xml_declaration=True)
                    continue

                if settings.clean_chart_labels and name.startswith("word/charts/") and name.endswith(".xml"):
                    root = _safe_fromstring(data)
                    if _clean_chart_labels(root):
                        updates[name] = etree.tostring(root, encoding="UTF-8", xml_declaration=True)
                    continue

                if settings.clean_language_settings and name.startswith("word/") and name.endswith(".xml"):
                    try:
                        root = _safe_fromstring(data)
                    except Exception:
                        continue
                    if _remove_lang_elements(root):
                        updates[name] = etree.tostring(root, encoding="UTF-8", xml_declaration=True)
                    continue

                if settings.clean_form_defaults and name.startswith("word/") and name.endswith(".xml"):
                    try:
                        root = _safe_fromstring(data)
                    except Exception:
                        continue
                    if _strip_form_defaults(root):
                        updates[name] = etree.tostring(root, encoding="UTF-8", xml_declaration=True)

        if not updates and not removed_parts:
            return

        temp_path = path + ".tmp_harden"
        with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(temp_path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for idx, item in enumerate(zin.infolist()):
                name = item.filename
                if name in removed_parts:
                    continue
                if idx != last_index.get(name, idx):
                    continue
                data = updates.get(name, zin.read(name))
                zout.writestr(item, data)

        os.replace(temp_path, path)

    def _iter_part_elements(self) -> Iterable[Any]:
        if self.doc.element is not None:
            yield self.doc.element
        for rel in self.doc.part.rels.values():
            if any(key in rel.reltype for key in ("header", "footer", "footnotes", "endnotes")):
                if hasattr(rel.target_part, "element"):
                    yield rel.target_part.element

    def _unlink_hyperlinks(self) -> None:
        for root in self._iter_part_elements():
            for hyperlink in list(root.iter(qn("w:hyperlink"))):
                parent = hyperlink.getparent()
                if parent is None:
                    continue
                idx = parent.index(hyperlink)
                for child in list(hyperlink):
                    parent.insert(idx, child)
                    idx += 1
                parent.remove(hyperlink)

    def _strip_comment_markers(self) -> None:
        tags = [qn("w:commentRangeStart"), qn("w:commentRangeEnd"), qn("w:commentReference")]
        for root in self._iter_part_elements():
            for tag in tags:
                for el in list(root.iter(tag)):
                    parent = el.getparent()
                    if parent is not None:
                        parent.remove(el)

    def _parse_merge_field_name(self, instr_text: str) -> Optional[str]:
        if not instr_text:
            return None
        normalized = " ".join(instr_text.replace("\u00a0", " ").split())
        if "MERGEFIELD" not in normalized.upper():
            return None
        parts = re.split(r"MERGEFIELD", normalized, flags=re.IGNORECASE, maxsplit=1)
        if len(parts) < 2:
            return None
        remainder = parts[1].strip()
        if remainder.startswith('"'):
            end_quote = remainder.find('"', 1)
            if end_quote != -1:
                return remainder[1:end_quote].strip()
        name = remainder.split("\\", 1)[0].strip()
        name = name.strip('"')
        return name or None

    def _build_merge_field_run(self, text: str, template_run=None):
        run = OxmlElement("w:r")
        if template_run is not None:
            rpr = template_run.find(qn("w:rPr"))
            if rpr is not None:
                run.append(copy.deepcopy(rpr))
        t = OxmlElement("w:t")
        t.text = text
        run.append(t)
        return run

    def _convert_mail_merge_fields(self) -> None:
        def _replace_simple_fields(root):
            for fld in list(root.iter(qn("w:fldSimple"))):
                instr = fld.get(qn("w:instr")) or ""
                name = self._parse_merge_field_name(instr)
                if not name:
                    continue
                template_run = None
                for child in fld.iter(qn("w:r")):
                    template_run = child
                    break
                new_run = self._build_merge_field_run(f"<<{name}>>", template_run)
                parent = fld.getparent()
                if parent is None:
                    continue
                idx = parent.index(fld)
                parent.remove(fld)
                parent.insert(idx, new_run)

        def _replace_complex_fields(root):
            for para in list(root.iter(qn("w:p"))):
                runs = list(para.findall(qn("w:r")))
                i = 0
                while i < len(runs):
                    run = runs[i]
                    fld_char = run.find(qn("w:fldChar"))
                    if fld_char is None or fld_char.get(qn("w:fldCharType")) != "begin":
                        i += 1
                        continue
                    instr_text = ""
                    display_text = ""
                    template_run = None
                    end_idx = None
                    in_result = False
                    j = i + 1
                    while j < len(runs):
                        current = runs[j]
                        instr_el = current.find(qn("w:instrText"))
                        if instr_el is not None and instr_el.text:
                            instr_text += instr_el.text
                        fld_mid = current.find(qn("w:fldChar"))
                        if fld_mid is not None:
                            fld_type = fld_mid.get(qn("w:fldCharType"))
                            if fld_type == "separate":
                                in_result = True
                            elif fld_type == "end":
                                end_idx = j
                                break
                        if in_result:
                            t_el = current.find(qn("w:t"))
                            if t_el is not None and t_el.text:
                                display_text += t_el.text
                                if template_run is None:
                                    template_run = current
                        j += 1

                    name = self._parse_merge_field_name(instr_text)
                    if name is None and display_text:
                        trimmed = display_text.strip()
                        trimmed = trimmed.replace("\u00ab", "").replace("\u00bb", "")
                        trimmed = trimmed.strip("<>").strip()
                        name = trimmed if trimmed else None

                    if name and end_idx is not None:
                        new_run = self._build_merge_field_run(f"<<{name}>>", template_run or run)
                        for k in range(end_idx, i - 1, -1):
                            para.remove(runs[k])
                        para.insert(i, new_run)
                        runs = list(para.findall(qn("w:r")))
                        i += 1
                        continue
                    i += 1

        for root in self._iter_part_elements():
            _replace_simple_fields(root)
            _replace_complex_fields(root)

    def _remove_data_bindings(self) -> None:
        for root in self._iter_part_elements():
            for sdt_pr in list(root.iter(qn("w:sdtPr"))):
                for binding in list(sdt_pr.iter(qn("w:dataBinding"))):
                    parent = binding.getparent()
                    if parent is not None:
                        parent.remove(binding)

    def _remove_hidden_text(self) -> None:
        hidden_tags = [qn("w:vanish"), qn("w:specVanish"), qn("w:webHidden")]
        for root in self._iter_part_elements():
            for run in list(root.iter(qn("w:r"))):
                rpr = run.find(qn("w:rPr"))
                if rpr is None:
                    continue
                if any(rpr.find(tag) is not None for tag in hidden_tags):
                    parent = run.getparent()
                    if parent is not None:
                        parent.remove(run)

    def _remove_invisible_objects(self) -> None:
        for root in self._iter_part_elements():
            for el in list(root.iter()):
                style = (el.get("style") or "").lower()
                visibility = (el.get("visibility") or "").lower()
                display = (el.get("display") or "").lower()
                if ("visibility:hidden" in style
                        or "visibility: hidden" in style
                        or "display:none" in style
                        or "mso-hide:all" in style
                        or visibility == "hidden"
                        or display == "none"):
                    parent = el.getparent()
                    if parent is not None:
                        parent.remove(el)

    def _remove_headers_footers(self) -> None:
        try:
            for sect_pr in list(self.doc.element.iter(qn("w:sectPr"))):
                for tag in (qn("w:headerReference"), qn("w:footerReference")):
                    for el in list(sect_pr.iter(tag)):
                        parent = el.getparent()
                        if parent is not None:
                            parent.remove(el)
        except Exception:
            pass

        try:
            rels = self.doc.part.rels
            for r_id, rel in list(rels.items()):
                if "header" in rel.reltype or "footer" in rel.reltype:
                    del rels[r_id]
        except Exception:
            pass

    def _remove_watermarks(self) -> None:
        def _is_watermark_candidate(el) -> bool:
            attrs = " ".join(str(v) for v in el.attrib.values()).lower()
            if "powerpluswatermarkobject" in attrs or "watermark" in attrs:
                return True
            if "mso-position-horizontal:center" in attrs and "mso-position-vertical:center" in attrs and "z-index" in attrs:
                return True
            return False

        for rel in self.doc.part.rels.values():
            if "header" not in rel.reltype:
                continue
            if not hasattr(rel.target_part, "element"):
                continue
            root = rel.target_part.element
            for el in list(root.iter()):
                if _is_watermark_candidate(el):
                    parent = el.getparent()
                    if parent is not None:
                        parent.remove(el)

    def _remove_ink_annotations(self) -> None:
        for root in self._iter_part_elements():
            for el in list(root.iter()):
                tag = el.tag
                if not isinstance(tag, str):
                    continue
                local = tag.split("}", 1)[-1].lower()
                if local.startswith("ink"):
                    parent = el.getparent()
                    if parent is not None:
                        parent.remove(el)

    def _scan_drawing_tag(self, drawing_element):
        """Deep traversal for Text Boxes inside w:drawing or w:pict."""
        from docx.text.paragraph import Paragraph
        from docx.oxml.ns import qn
        
        # w:drawing -> ... -> w:txbxContent -> w:p
        for txbx in drawing_element.iter(qn('w:txbxContent')):
            for content_node in txbx:
                if content_node.tag == qn('w:p'):
                    p = Paragraph(content_node, self.doc)
                    self._scan_paragraph(p)
                elif content_node.tag == qn('w:tbl'):
                     # recursive scan for table inside text box
                     self._scan_table_xml(content_node)

    def _scan_run_contents(self, run_element):
        """Scan a w:r element for nested drawings."""
        from docx.oxml.ns import qn
        # w:r -> w:drawing | w:pict
        for tag_name in ['w:drawing', 'w:pict']:
             for drawing in run_element.iter(qn(tag_name)):
                 self._scan_drawing_tag(drawing)

    def _append_run(self, para, run):
        t = run.text or ""
        for i, ch in enumerate(t):
            self.text += ch
            self.index.append((para, run, i))

    def _scan_paragraph(self, para):
        from docx.text.run import Run
        from docx.text.paragraph import Paragraph
        from docx.oxml.ns import qn
        
        # Iterate over children to catch runs inside hyperlinks AND drawings (text boxes)
        for child in para._element:
            if child.tag == qn('w:r'):
                run = Run(child, para)
                self._append_run(para, run)
                self._scan_run_contents(child)
            elif child.tag == qn('w:hyperlink'):
                for subchild in child:
                    if subchild.tag == qn('w:r'):
                        run = Run(subchild, para)
                        self._append_run(para, run)
                        self._scan_run_contents(subchild)
                    elif subchild.tag == qn('w:ins'):
                        for node in subchild:
                            if node.tag == qn('w:r'):
                                run = Run(node, para)
                                self._append_run(para, run)
                                self._scan_run_contents(node)
            elif child.tag == qn('w:drawing') or child.tag == qn('w:pict'):
                # Direct child drawings (less common but supported)
                self._scan_drawing_tag(child)
            elif child.tag == qn('w:fldSimple'):
                # Simple Field (scan runs inside)
                for subchild in child:
                    if subchild.tag == qn('w:r'):
                        run = Run(subchild, para)
                        self._append_run(para, run)
                        self._scan_run_contents(subchild)
                    elif subchild.tag == qn('w:ins'):
                        for node in subchild:
                            if node.tag == qn('w:r'):
                                run = Run(node, para)
                                self._append_run(para, run)
                                self._scan_run_contents(node)
            elif child.tag == qn('w:sdt'):
                # Inline Content Control
                # w:sdt -> w:sdtContent -> w:r
                for sdt_content in child.iter(qn('w:sdtContent')):
                    for subchild in sdt_content:
                        if subchild.tag == qn('w:r'):
                            run = Run(subchild, para)
                            self._append_run(para, run)
                            self._scan_run_contents(subchild)
                        elif subchild.tag == qn('w:ins'):
                            for node in subchild:
                                if node.tag == qn('w:r'):
                                    run = Run(node, para)
                                    self._append_run(para, run)
                                    self._scan_run_contents(node)
                        # Could recursively support other inline content here
            elif child.tag == qn('w:ins'):
                # Tracked Changes Insertion -> w:r
                for subchild in child:
                    if subchild.tag == qn('w:r'):
                        run = Run(subchild, para)
                        self._append_run(para, run)
                        self._scan_run_contents(subchild)
        
        self.text += "\n"; self.index.append(("break", None, None))

    def _scan_container(self, container):
        from docx.text.paragraph import Paragraph
        from docx.oxml.ns import qn
        
        # Robustly get XML element to iterate
        element = getattr(container, 'element', getattr(container, '_element', None))
        if element is None:
            return

        # Handle Document object which wraps w:document -> w:body
        if element.tag == qn('w:document'):
            element = element.body
            
        # Iterate all children to catch w:sdt (Content Controls)
        for child in element:
            if child.tag == qn('w:p'):
                p = Paragraph(child, self.doc)
                self._scan_paragraph(p)
            elif child.tag == qn('w:tbl'):
                self._scan_table_xml(child)
            elif child.tag == qn('w:sdt'):
                # Block Level Content Control
                # w:sdt -> w:sdtContent -> (w:p | w:tbl | w:sdt)
                for sdt_content in child.iter(qn('w:sdtContent')):
                    for node in sdt_content:
                        if node.tag == qn('w:p'):
                            p = Paragraph(node, self.doc)
                            self._scan_paragraph(p)
                        elif node.tag == qn('w:tbl'):
                            self._scan_table_xml(node)
                        # We could recurse for nested SDTs, but 1 level deep often suffices.
                        # For true recursion we'd need a helper, but this covers 99%.

    def _scan_table_xml(self, tbl_element):
        """Recursively scan a table element (w:tbl) found inside another container."""
        from docx.text.paragraph import Paragraph
        from docx.oxml.ns import qn
        
        # w:tbl -> w:tr -> w:tc -> (w:p | w:tbl)
        for tr in tbl_element.iter(qn('w:tr')):
            for tc in tr.iter(qn('w:tc')):
                for child in tc:
                    if child.tag == qn('w:p'):
                        p = Paragraph(child, self.doc)
                        self._scan_paragraph(p)
                    elif child.tag == qn('w:tbl'):
                        # Nested table inside a cell
                        self._scan_table_xml(child)
                    elif child.tag == qn('w:sdt'):
                        # Sdt inside table cell?
                         # w:sdt -> w:sdtContent -> (w:p | w:r)
                        for sdt_content in child.iter(qn('w:sdtContent')):
                            for node in sdt_content:
                                if node.tag == qn('w:p'):
                                    p = Paragraph(node, self.doc)
                                    self._scan_paragraph(p)
                                elif node.tag == qn('w:r'):
                                    # Scan runs inside table/SDT for drawings too
                                    self._scan_run_contents(node)


    def _build(self):
        # Scan headers and footers first (to ensure they are indexed)
        # Use set of ELEMENT IDs to avoid duplicates (proxy objects have different IDs)
        scanned_elements = set()

        for section in self.doc.sections:
            # Check Headers
            for header in [section.header, section.first_page_header, section.even_page_header]:
                if header and not header.is_linked_to_previous:
                    # Deduplicate based on the XML element ID
                    el_id = id(header._element)
                    if el_id not in scanned_elements:
                        self._scan_container(header)
                        scanned_elements.add(el_id)
            
            # Check Footers
            for footer in [section.footer, section.first_page_footer, section.even_page_footer]:
                if footer and not footer.is_linked_to_previous:
                    el_id = id(footer._element)
                    if el_id not in scanned_elements:
                        self._scan_container(footer)
                        scanned_elements.add(el_id)

        # Scan footnotes
        try:
             part = self.doc.part
             footnote_part = None
             for rel in part.rels.values():
                 if "footnotes" in rel.reltype:
                     footnote_part = rel.target_part
                     break
             
             if footnote_part:
                 from docx.text.paragraph import Paragraph
                 from docx.oxml.ns import qn
                 from docx.oxml import parse_xml

                 # Check if we have a live element or need to parse from blob
                 root = None
                 try:
                     if hasattr(footnote_part, 'element'):
                         root = footnote_part.element
                 except Exception:
                     # Accessing .element failed (likely generic Part)
                     root = None

                 if root is None:
                     root = parse_xml(footnote_part.blob)
                     # Register for saving later
                     self.detached_parts.append((footnote_part, root))
                 
                 # iterate w:footnote elements
                 
                 # iterate w:footnote elements
                 for fn in root.findall(qn('w:footnote')):
                     for child in fn:
                         if child.tag == qn('w:p'):
                             # Wrap and scan
                             p = Paragraph(child, self.doc)
                             self._scan_paragraph(p)
                         elif child.tag == qn('w:tbl'):
                             self._scan_table_xml(child)
                         elif child.tag == qn('w:sdt'):
                             # SDT in Footnote
                             for sdt_content in child.iter(qn('w:sdtContent')):
                                 for node in sdt_content:
                                     if node.tag == qn('w:p'):
                                         p = Paragraph(node, self.doc)
                                         self._scan_paragraph(p)
                                     elif node.tag == qn('w:tbl'):
                                         self._scan_table_xml(node)
        except Exception:
            # Safely ignore footnote errors to prevent crash
            pass

        # Scan endnotes
        try:
             part = self.doc.part
             endnote_part = None
             for rel in part.rels.values():
                 if "endnotes" in rel.reltype:
                     endnote_part = rel.target_part
                     break
             
             if endnote_part:
                 from docx.text.paragraph import Paragraph
                 from docx.oxml.ns import qn
                 from docx.oxml import parse_xml

                 # Robustly get the root element
                 root = None
                 try: 
                     if hasattr(endnote_part, 'element'):
                         root = endnote_part.element
                 except Exception:
                     root = None
                 
                 if root is None:
                     root = parse_xml(endnote_part.blob)
                     # Register for saving later
                     self.detached_parts.append((endnote_part, root))
                 
                 # iterate w:endnote elements
                 for en in root.findall(qn('w:endnote')):
                     for child in en:
                         if child.tag == qn('w:p'):
                             # Wrap and scan
                             p = Paragraph(child, self.doc)
                             self._scan_paragraph(p)
                         elif child.tag == qn('w:tbl'):
                             self._scan_table_xml(child)
                         elif child.tag == qn('w:sdt'):
                             # SDT in Endnote
                             for sdt_content in child.iter(qn('w:sdtContent')):
                                 for node in sdt_content:
                                     if node.tag == qn('w:p'):
                                         p = Paragraph(node, self.doc)
                                         self._scan_paragraph(p)
                                     elif node.tag == qn('w:tbl'):
                                         self._scan_table_xml(node)
        except Exception:
            pass

        # Scan main document body
        self._scan_container(self.doc)

    def _make_text_run(self, text: str, rPr=None, source_run=None) -> OxmlElement:
        r = OxmlElement("w:r")
        if source_run is not None:
            # Copy attributes (like rsidR) from source to preserve identity
            for key, value in source_run.attrib.items():
                r.set(key, value)
        
        if rPr is not None:
            from copy import deepcopy
            r.append(deepcopy(rPr))
        
        t = OxmlElement("w:t")
        # Preserve spaces so replacements aren't collapsed
        t.set(qn('xml:space'), 'preserve')
        t.text = text
        r.append(t)
        return r

    def _insert_after(self, run: Run, text: str, highlight: bool) -> Run:
        # Legacy highlighter insertion; kept for compatibility
        r = self._make_text_run(text)
        run._element.addnext(r)
        new = Run(r, run._parent)
        if highlight:
            new.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return new

    def _insert_deletion_after(self, anchor_el, text: str, rPr=None):
        del_el = OxmlElement('w:del')
        del_el.set(qn('w:id'), str(self._rev_id)); self._rev_id += 1
        del_el.set(qn('w:author'), self.author_name)
        del_el.set(qn('w:date'), datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ'))
        
        # Create run for deleted text, retaining original formatting
        r = self._make_text_run(text, rPr)
        
        # Change w:t to w:delText
        t_el = r.find(qn('w:t'))
        if t_el is not None:
            t_el.tag = qn('w:delText')
            
        del_el.append(r)
        anchor_el.addnext(del_el)
        return del_el

    def _insert_insertion_after(self, anchor_el, text: str, rPr=None):
        ins_el = OxmlElement('w:ins')
        ins_el.set(qn('w:id'), str(self._rev_id)); self._rev_id += 1
        ins_el.set(qn('w:author'), self.author_name)
        ins_el.set(qn('w:date'), datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ'))
        r = self._make_text_run(text, rPr)
        ins_el.append(r)
        anchor_el.addnext(ins_el)
        return ins_el



    def harden_document(
        self,
        scrub_all_images: bool = False,
        settings: Optional[MetadataCleaningSettings] = None
    ):
        """
        Apply security hardening measures:
        1. Remove Hyperlinks (convert to plain text by deleting relationships).
        2. Clear RSIDs (Revision Save IDs) to prevent fingerprinting.
        3. Replace Embedded Objects with placeholders.
        4. (Optional) Delete all images if enabled.
        5. (Implicit) Thumbnail deletion is best done on file save/zip level, but we can try removing the package rel.
        """
        if settings is None:
            settings = MetadataCleaningSettings()

        # 1. Remove Hyperlinks (delete relationships to convert links to plain text)
        if settings.clean_hyperlink_urls:
            try:
                self._unlink_hyperlinks()
                # Collect hyperlink relationship IDs to remove
                hyperlink_rel_ids = []
                for rel_id, rel in self.doc.part.rels.items():
                    if "hyperlink" in rel.reltype:
                        hyperlink_rel_ids.append(rel_id)

                # Remove the hyperlink relationships
                # This converts the hyperlinks to plain text in the document
                for rel_id in hyperlink_rel_ids:
                    del self.doc.part.rels[rel_id]
            except Exception:
                pass

        # 2. Hardening Pass (RSID, Objects, Images)
        # We need to traverse Body, Headers, Footers, Footnotes, Endnotes
        
        def _harden_element(element):
            # A. Clear RSID attributes
            if settings.clean_rsids:
                for key in list(element.attrib.keys()):
                    if qn('w:rsidR') in key or qn('w:rsidRPr') in key or qn('w:rsidP') in key:
                        del element.attrib[key]

            # B. Check contents
            # We must iterate a copy of children to modify structure safely
            for child in list(element):
                tag = child.tag
                
                # Embedded Objects
                if tag == qn('w:object') or tag == qn('w:control'):
                    if settings.clean_ole_objects or settings.clean_activex:
                        parent = element
                        if parent.tag == qn('w:r'):
                            parent.remove(child)
                            if len(parent) == 0:
                                t_repl = OxmlElement('w:t')
                                t_repl.text = "[REDACTED DATA OBJECT]"
                                parent.append(t_repl)
                        else:
                            replacement_run = OxmlElement('w:r')
                            t_repl = OxmlElement('w:t')
                            t_repl.text = "[REDACTED DATA OBJECT]"
                            replacement_run.append(t_repl)
                            index = parent.index(child)
                            parent.insert(index, replacement_run)
                            parent.remove(child)
                        continue

                # Images (Drawings)
                if scrub_all_images:
                    if tag == qn('w:drawing') or tag == qn('w:pict'):
                         # Delete image
                         element.remove(child)
                         continue
                         
                # Recurse
                _harden_element(child)

        # Apply to Body
        if self.doc.element.body is not None:
            _harden_element(self.doc.element.body)

        # Apply to Headers/Footers
        try:
             for rel in self.doc.part.rels.values():
                if "header" in rel.reltype or "footer" in rel.reltype:
                    if hasattr(rel.target_part, 'element'):
                       _harden_element(rel.target_part.element)
        except Exception:
            pass

    def scrub_metadata(self, settings: Optional[MetadataCleaningSettings] = None):
        """
        Wipes document metadata based on user-configurable settings.
        If no settings provided, uses defaults (all ON except created/modified dates).
        Values are deleted (set to empty) rather than replaced with placeholder text.
        """
        if settings is None:
            settings = MetadataCleaningSettings()
        
        cp = self.doc.core_properties
        
        # Core Properties - delete values (use empty strings)
        if settings.clean_author:
            cp.author = ""
        if settings.clean_last_modified_by:
            cp.last_modified_by = ""
        if settings.clean_comments:
            cp.comments = ""
        if settings.clean_title:
            cp.title = ""
        if settings.clean_subject:
            cp.subject = ""
        if settings.clean_keywords:
            cp.keywords = ""
        if settings.clean_category:
            try:
                cp.category = ""
            except Exception:
                pass
        if settings.clean_content_status:
            try:
                cp.content_status = ""
            except Exception:
                pass
        if settings.clean_revision_number:
            cp.revision = 1
        if settings.clean_identifier:
            try:
                cp.identifier = ""
            except Exception:
                pass
        if settings.clean_language:
            try:
                cp.language = ""
            except Exception:
                pass
        if settings.clean_version:
            try:
                cp.version = ""
            except Exception:
                pass
        if settings.clean_last_printed:
            try:
                cp.last_printed = None
            except Exception:
                pass
        if settings.clean_created_date:
            try:
                cp.created = None
            except Exception:
                pass
        if settings.clean_modified_date:
            try:
                cp.modified = None
            except Exception:
                pass

        # Remove empty core property elements to keep schema-valid content
        try:
            from lxml import etree

            core_part = None
            for rel in self.doc.part.package.rels.values():
                if "core-properties" in rel.reltype:
                    core_part = rel.target_part
                    break

            if core_part is not None:
                root = core_part.element if hasattr(core_part, "element") else etree.fromstring(core_part._blob)
                ns = {
                    "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
                    "dc": "http://purl.org/dc/elements/1.1/",
                    "dcterms": "http://purl.org/dc/terms/",
                }

                def _remove(tag: str) -> None:
                    for elem in root.findall(tag, namespaces=ns):
                        parent = elem.getparent()
                        if parent is not None:
                            parent.remove(elem)

                if settings.clean_author:
                    _remove("dc:creator")
                if settings.clean_last_modified_by:
                    _remove("cp:lastModifiedBy")
                if settings.clean_title:
                    _remove("dc:title")
                if settings.clean_subject:
                    _remove("dc:subject")
                if settings.clean_keywords:
                    _remove("cp:keywords")
                if settings.clean_comments:
                    _remove("dc:description")
                if settings.clean_category:
                    _remove("cp:category")
                if settings.clean_content_status:
                    _remove("cp:contentStatus")
                if settings.clean_identifier:
                    _remove("dc:identifier")
                if settings.clean_language:
                    _remove("dc:language")
                if settings.clean_version:
                    _remove("cp:version")
                if settings.clean_last_printed:
                    _remove("cp:lastPrinted")
                if settings.clean_created_date:
                    _remove("dcterms:created")
                if settings.clean_modified_date:
                    _remove("dcterms:modified")

                if settings.clean_revision_number:
                    for elem in root.findall("cp:revision", namespaces=ns):
                        elem.text = "1"

                core_part._blob = etree.tostring(root, encoding="UTF-8", xml_declaration=True)
        except Exception as e:
            print(f"[MARCUT] Warning: core.xml cleaning failed: {e}")

        # 1. Disable Spell Check/Grammar
        if settings.clean_spell_grammar_state:
            try:
                doc_settings = self.doc.settings.element
                proof_states = list(doc_settings.findall(qn('w:proofState')))
                if proof_states:
                    # Keep the first proofState, normalize attributes, drop extras
                    primary = proof_states[0]
                    primary.set(qn('w:spelling'), 'clean')
                    primary.set(qn('w:grammar'), 'clean')
                    for attr in list(primary.attrib.keys()):
                        if attr not in (qn('w:spelling'), qn('w:grammar')):
                            del primary.attrib[attr]
                    for extra in proof_states[1:]:
                        doc_settings.remove(extra)
                else:
                    proof_state = OxmlElement('w:proofState')
                    proof_state.set(qn('w:spelling'), 'clean')
                    proof_state.set(qn('w:grammar'), 'clean')
                    doc_settings.append(proof_state)
            except Exception:
                pass

        # 2. Remove Comments and Custom Properties parts via Relationships
        try:
            rels = self.doc.part.rels
            ids_to_remove = []
            
            for r_id, rel in rels.items():
                should_remove = False
                
                # Comments
                if settings.clean_review_comments and "comments" in rel.reltype:
                    should_remove = True
                # Custom XML (custom properties store here)
                if settings.clean_custom_properties and ("customXml" in rel.reltype or "custom-properties" in rel.reltype):
                    should_remove = True
                # Glossary
                if settings.clean_glossary and "glossary" in rel.reltype:
                    should_remove = True
                # VBA Macros
                if settings.clean_vba_macros and "vbaProject" in rel.reltype:
                    should_remove = True
                # Digital Signatures
                if settings.clean_digital_signatures and "signature" in rel.reltype:
                    should_remove = True
                
                if should_remove:
                    ids_to_remove.append(r_id)
                    try:
                        if hasattr(rel.target_part, 'element'):
                            rel.target_part.element.clear()
                    except Exception:
                        pass
            
            for r_id in ids_to_remove:
                del rels[r_id]
        except Exception:
            pass

        if settings.clean_review_comments:
            try:
                self._strip_comment_markers()
            except Exception:
                pass
        
        # 3. Clean app.xml properties (Company, Manager, Application, etc.)
        # These require direct XML access via the package parts
        try:
            from lxml import etree
            
            # Access the app.xml part through document relationships
            for rel in self.doc.part.package.rels.values():
                if "extended-properties" in rel.reltype or "app" in rel.reltype:
                    app_part = rel.target_part
                    if hasattr(app_part, '_blob') and app_part._blob:
                        # Parse the app.xml content
                        app_xml = etree.fromstring(app_part._blob)
                        
                        # Define namespace
                        ns = {'ep': 'http://schemas.openxmlformats.org/officeDocument/2006/extended-properties'}

                        def _remove_elements(tag: str) -> None:
                            for elem in app_xml.findall(f'.//ep:{tag}', namespaces=ns):
                                parent = elem.getparent()
                                if parent is not None:
                                    parent.remove(elem)
                        
                        # Clean Company
                        if settings.clean_company:
                            _remove_elements("Company")
                        
                        # Clean Manager
                        if settings.clean_manager:
                            _remove_elements("Manager")
                        
                        # Clean Application
                        if settings.clean_application:
                            _remove_elements("Application")
                        
                        # Clean AppVersion
                        if settings.clean_app_version:
                            _remove_elements("AppVersion")
                        
                        # Clean Template
                        if settings.clean_template:
                            _remove_elements("Template")

                        if settings.clean_hyperlink_base:
                            _remove_elements("HyperlinkBase")
                        
                        # Clean TotalTime (editing time)
                        if settings.clean_total_editing_time:
                            for elem in app_xml.findall('.//ep:TotalTime', namespaces=ns):
                                elem.text = "0"
                        
                        # Clean Statistics (Words, Characters, Lines, Paragraphs, Pages)
                        if settings.clean_statistics:
                            for tag in ['Words', 'Characters', 'CharactersWithSpaces', 'Lines', 'Paragraphs', 'Pages']:
                                for elem in app_xml.findall(f'.//ep:{tag}', namespaces=ns):
                                    elem.text = "0"
                        
                        # Clean DocSecurity
                        if settings.clean_doc_security:
                            for elem in app_xml.findall('.//ep:DocSecurity', namespaces=ns):
                                elem.text = "0"
                        
                        # Clean ScaleCrop
                        if settings.clean_scale_crop:
                            for elem in app_xml.findall('.//ep:ScaleCrop', namespaces=ns):
                                elem.text = "false"
                        
                        # Clean LinksUpToDate
                        if settings.clean_links_up_to_date:
                            for elem in app_xml.findall('.//ep:LinksUpToDate', namespaces=ns):
                                elem.text = "false"
                        
                        # Clean SharedDoc
                        if settings.clean_shared_doc:
                            for elem in app_xml.findall('.//ep:SharedDoc', namespaces=ns):
                                elem.text = "false"
                        
                        # Clean HyperlinksChanged
                        if settings.clean_hyperlinks_changed:
                            for elem in app_xml.findall('.//ep:HyperlinksChanged', namespaces=ns):
                                elem.text = "false"
                        
                        # Write back the modified XML
                        app_part._blob = etree.tostring(app_xml, encoding='UTF-8', xml_declaration=True)
                        break
        except Exception as e:
            # Log but don't fail - app.xml cleaning is best-effort
            print(f"[MARCUT] Warning: app.xml cleaning failed: {e}")

        # 4. Clean document settings (GUID, docVars, language, forms)
        try:
            doc_settings = self.doc.settings.element
            if settings.clean_document_guid:
                for el in list(doc_settings.iter(qn('w14:docId'))):
                    parent = el.getparent()
                    if parent is not None:
                        parent.remove(el)
            if settings.clean_document_variables:
                doc_vars = doc_settings.find(qn('w:docVars'))
                if doc_vars is not None:
                    doc_settings.remove(doc_vars)
            if settings.clean_fast_save_data:
                for tag in (qn('w:savePreviewPicture'), qn('w:saveThroughXslt')):
                    el = doc_settings.find(tag)
                    if el is not None:
                        doc_settings.remove(el)
            if settings.clean_mail_merge:
                mail_merge = doc_settings.find(qn('w:mailMerge'))
                if mail_merge is not None:
                    doc_settings.remove(mail_merge)
        except Exception:
            pass

        if settings.clean_mail_merge:
            try:
                self._convert_mail_merge_fields()
            except Exception:
                pass

        if settings.clean_data_bindings:
            try:
                self._remove_data_bindings()
            except Exception:
                pass

        if settings.clean_hidden_text:
            try:
                self._remove_hidden_text()
            except Exception:
                pass

        if settings.clean_invisible_objects:
            try:
                self._remove_invisible_objects()
            except Exception:
                pass

        if settings.clean_ink_annotations:
            try:
                self._remove_ink_annotations()
            except Exception:
                pass

        if settings.clean_watermarks:
            try:
                self._remove_watermarks()
            except Exception:
                pass

        if settings.clean_headers_footers:
            try:
                self._remove_headers_footers()
            except Exception:
                pass

        if settings.clean_alt_text:
            try:
                for root in self._iter_part_elements():
                    for el in root.iter():
                        if el.tag == '{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr':
                            if 'descr' in el.attrib:
                                el.attrib['descr'] = ""
                            if 'title' in el.attrib:
                                el.attrib['title'] = ""
            except Exception:
                pass

        if settings.clean_language_settings:
            try:
                for root in self._iter_part_elements():
                    for el in list(root.iter(qn('w:lang'))):
                        parent = el.getparent()
                        if parent is not None:
                            parent.remove(el)
            except Exception:
                pass

        if settings.clean_form_defaults:
            try:
                for root in self._iter_part_elements():
                    for tag in (qn('w:default'), qn('w:result')):
                        for el in list(root.iter(tag)):
                            parent = el.getparent()
                            if parent is not None:
                                parent.remove(el)
            except Exception:
                pass

        if settings.clean_hyperlink_urls:
            try:
                self._unlink_hyperlinks()
            except Exception:
                pass

        # Store settings for potential later use
        self._metadata_settings = settings



    def apply_replacements(self, spans: List[Dict[str,Any]], track_changes: bool = True):
        spans = sorted(spans, key=lambda s: s["start"], reverse=True)
        for sp in spans:
            s, e, repl = sp["start"], sp["end"], sp["replacement"]
            if not (0 <= s < len(self.index)):
                continue
            start_idx = self.index[s]
            if start_idx[0] == "break":
                continue
            para0, run0, ci0 = start_idx

            # Collect affected characters per run
            buckets: Dict[tuple, List] = {}
            for pos in range(s, min(e, len(self.index))):
                idx = self.index[pos]
                if idx[0] == "break":
                    continue
                key = (id(idx[0]), id(idx[1]))
                buckets.setdefault(key, []).append(idx)

            if not track_changes:
                # Legacy: remove and optionally highlight insert
                for _, chars in buckets.items():
                    r = chars[0][1]
                    buf = list(r.text or "")
                    for _, _, ci in chars:
                        if 0 <= ci < len(buf):
                            buf[ci] = ""
                    r.text = "".join(buf)
                before = run0.text[:ci0]
                after = run0.text[ci0:]
                run0.text = before
                tag = self._insert_after(run0, repl, False)
                self._insert_after(tag, after, False)
                continue

            # Track changes mode: wrap deletions and add insertion
            # First, process each run that intersects with the span
            processed_first = False
            for _, chars in buckets.items():
                r = chars[0][1]
                original = r.text or ""
                cis = sorted(ci for _, _, ci in chars if ci is not None)
                if not cis:
                    continue
                start_ci = cis[0]; end_ci = cis[-1]
                pre = original[:start_ci]
                mid = original[start_ci:end_ci+1]
                post = original[end_ci+1:]

                # Grab the run properties to preserve formatting
                rPr = r._element.find(qn('w:rPr'))

                # Set the run to pre-text only
                r.text = pre
                
                # Insert deletion for removed text (preserving formatting)
                del_el = self._insert_deletion_after(r._element, mid, rPr)

                if not processed_first and r is run0:
                    # Insert the replacement as an insertion after the deletion
                    # Preserve the original run's formatting (font, size, bold, italic) so that
                    # redaction labels match the surrounding text. We pass rPr and then strip
                    # only problematic properties (hidden, shading) while keeping visual styling.
                    # ALSO: Force RED color to ensure visibility against any background.
                    ins_el = self._insert_insertion_after(del_el, repl, rPr)
                    try:
                         # Find w:r inside w:ins
                         r_el = ins_el.find(qn('w:r'))
                         if r_el is not None:
                             rPr_el = r_el.find(qn('w:rPr'))
                             if rPr_el is None:
                                 rPr_el = OxmlElement('w:rPr')
                                 r_el.insert(0, rPr_el)
                             
                             # Strip problematic properties that could hide or obscure the label
                             # while preserving font, size, bold, italic, underline
                             STRIP_PROPS = ['w:vanish', 'w:webHidden', 'w:shd', 'w:highlight', 
                                           'w:effect', 'w:specVanish', 'w:oMath']
                             for prop in STRIP_PROPS:
                                 el = rPr_el.find(qn(prop))
                                 if el is not None:
                                     rPr_el.remove(el)
                             
                             # Force Color Red for visibility
                             # First remove any existing color to avoid duplicates
                             existing_color = rPr_el.find(qn('w:color'))
                             if existing_color is not None:
                                 rPr_el.remove(existing_color)
                             color_el = OxmlElement('w:color')
                             color_el.set(qn('w:val'), 'FF0000')
                             rPr_el.append(color_el)
                    except Exception:
                        pass
                    
                    # Append the remainder of this run after the insertion
                    if post:
                        # CRITICAL FIX: Use rPr for the post-split run so it doesn't lose formatting
                        # AND pass source run element to preserve attributes (rsid) so it's not seen as new insert
                        ins_el.addnext(self._make_text_run(post, rPr, r._element))
                    processed_first = True
                else:
                    # For other runs, just append the remainder after deletion
                    if post:
                        # CRITICAL FIX: Use rPr for the post-split run
                        del_el.addnext(self._make_text_run(post, rPr, r._element))
