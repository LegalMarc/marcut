"""Programmatic DOCX fixture builders for the docx_io golden-file harness
(issue #70 / docs/design/docx_io_package_split.md §3.2 item 1).

Every fixture is generated from scratch at test time (no binary blobs
committed to the repo) so the harness stays reviewable and has no stale
binary drift. Determinism is the load-bearing property here: two runs of
`build_kitchen_sink_docx()` on the same interpreter must produce byte-for-
byte identical `.docx` output, because the golden JSON snapshots in
`tests/docx_io_golden/golden/` are diffed against a fresh build every time
the harness runs. In particular:

- `core_properties.created`/`.modified` are pinned to a fixed UTC instant
  (python-docx would otherwise leave whatever the template shipped with,
  which is fine, but pinning removes any doubt).
- Every author/date on a *pre-existing* tracked-change element (`w:ins`/
  `w:del` already present "on load") is a fixed literal, never
  `datetime.now()`.
- The one source of real non-determinism -- `DocxMap.apply_replacements()`
  stamping fresh `w:ins`/`w:del` elements with `datetime.now(timezone.utc)`
  -- is handled on the harness side (see `harness.py`'s
  `_normalize_dynamic_attrs`), not here.

The fixture covers every container type enumerated in the issue:
plain paragraphs; tables incl. a nested table; headers/footers; footnotes;
endnotes; a legacy VML textbox/drawing; block-level and inline content
controls (`w:sdt`); tracked-changes already present on load; an embedded
OLE object; a hyperlink; comments in all three visibility states the
codebase's `_comment_visibility_map()` distinguishes (visible/hidden/
deleted); and a mail-merge field.

Each piece of redactable body text carries a distinct, greppable marker
(`BODYMARK`, `TABLEMARK`, ...) so the harness can locate it in
`DocxMap.text` by substring search and build `apply_replacements()` spans
without hard-coding character offsets that would drift if this file
changes.
"""
from __future__ import annotations

import io
from datetime import datetime, timezone
from xml.etree import ElementTree as ET
import zipfile

from docx import Document

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
V_NS = "urn:schemas-microsoft-com:vml"
O_NS = "urn:schemas-microsoft-com:office:office"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"

FIXED_DATE = datetime(2020, 1, 1, tzinfo=timezone.utc)
PRIOR_REVISION_DATE = "2020-01-01T00:00:00Z"

# Markers are deliberately distinct substrings (no marker is a substring of
# another) so `DocxMap.text.index(marker)` can never match the wrong one.
MARKERS = {
    "body": "BODYMARK_a1",
    "table": "TABLEMARK_b2",
    "nested_table": "NESTEDTABLEMARK_c3",
    "header": "HEADERMARK_d4",
    "footer": "FOOTERMARK_e5",
    "footnote": "FOOTNOTEMARK_f6",
    "endnote": "ENDNOTEMARK_g7",
    "textbox": "TEXTBOXMARK_h8",
    "sdt_block": "SDTBLOCKMARK_i9",
    "sdt_inline": "SDTINLINEMARK_j10",
    "hyperlink": "HYPERLINKMARK_k11",
    "comment_visible_anchor": "COMMENTVISIBLEMARK_l12",
    "comment_hidden_anchor": "COMMENTHIDDENMARK_m13",
    "prior_insert": "PRIORINSERTMARK_n14",
}

# Markers that DocxMap._build() is expected to index and that the harness
# therefore includes in its apply_replacements() spans. Deliberately
# excludes markers that live in prose the scanner never sees (comment
# bodies, deleted-on-load text, mail-merge instructions) -- their absence
# from redaction is itself part of what the golden snapshot characterizes.
REDACTABLE_MARKERS = (
    "body",
    "table",
    "nested_table",
    "header",
    "footer",
    "footnote",
    "endnote",
    "textbox",
    "sdt_block",
    "sdt_inline",
    "hyperlink",
    "comment_visible_anchor",
    "comment_hidden_anchor",
    "prior_insert",
)


def _qn(ns: str, tag: str) -> str:
    return f"{{{ns}}}{tag}"


def _read_zip_entry(path: str, name: str) -> bytes:
    with zipfile.ZipFile(path) as zf:
        return zf.read(name)


def _patch_zip(path: str, updates: dict, new_entries: dict | None = None) -> None:
    new_entries = new_entries or {}
    buf = io.BytesIO()
    with zipfile.ZipFile(path, "r") as zin:
        existing = {i.filename for i in zin.infolist()}
        with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = updates.get(item.filename, zin.read(item.filename))
                # Pin the ZIP member timestamp so byte-identical fixture
                # rebuilds don't depend on wall-clock time at all (the
                # harness canonicalizes XML content rather than comparing
                # raw ZIP bytes, but keeping this deterministic too costs
                # nothing and makes ad-hoc `diff` of two fixture builds
                # meaningful).
                item.date_time = (2020, 1, 1, 0, 0, 0)
                zout.writestr(item, data)
            for name, data in new_entries.items():
                if name not in existing:
                    info = zipfile.ZipInfo(name, date_time=(2020, 1, 1, 0, 0, 0))
                    info.compress_type = zipfile.ZIP_DEFLATED
                    zout.writestr(info, data)
    with open(path, "wb") as f:
        f.write(buf.getvalue())


def _add_default_content_type(content_types_xml: bytes, extension: str, content_type: str) -> bytes:
    root = ET.fromstring(content_types_xml)
    for default in root.findall(_qn(CT_NS, "Default")):
        if default.get("Extension") == extension:
            return content_types_xml  # already registered, don't collide
    ET.SubElement(root, _qn(CT_NS, "Default"), {
        "Extension": extension,
        "ContentType": content_type,
    })
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _add_override_content_type(content_types_xml: bytes, part_name: str, content_type: str) -> bytes:
    root = ET.fromstring(content_types_xml)
    for override in root.findall(_qn(CT_NS, "Override")):
        if override.get("PartName") == part_name:
            return content_types_xml
    ET.SubElement(root, _qn(CT_NS, "Override"), {
        "PartName": part_name,
        "ContentType": content_type,
    })
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _add_relationship(rels_xml: bytes, rel_id: str, rel_type: str, target: str,
                       target_mode: str | None = None) -> bytes:
    root = ET.fromstring(rels_xml)
    for rel in root.findall(_qn(REL_NS, "Relationship")):
        if rel.get("Id") == rel_id:
            return rels_xml
    attrs = {"Id": rel_id, "Type": rel_type, "Target": target}
    if target_mode:
        attrs["TargetMode"] = target_mode
    ET.SubElement(root, _qn(REL_NS, "Relationship"), attrs)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def build_kitchen_sink_docx(path: str) -> None:
    """Write a single DOCX fixture at `path` covering every container type
    listed in issue #70. Deterministic across repeated calls."""
    doc = Document()
    cp = doc.core_properties
    cp.created = FIXED_DATE
    cp.modified = FIXED_DATE
    cp.author = "Golden Fixture Author"
    cp.title = "docx_io golden fixture"

    doc.add_paragraph("Plain paragraph with no PII-like content.")
    doc.add_paragraph(f"Body paragraph containing {MARKERS['body']} inline.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = f"Table cell with {MARKERS['table']}."
    nested_cell = table.cell(1, 1)
    nested_table = nested_cell.add_table(rows=1, cols=1)
    nested_table.cell(0, 0).text = f"Nested table cell with {MARKERS['nested_table']}."

    section = doc.sections[0]
    section.header.paragraphs[0].text = f"Header with {MARKERS['header']}."
    section.footer.paragraphs[0].text = f"Footer with {MARKERS['footer']}."

    doc.save(path)

    _inject_footnote_endnote(path)
    _inject_comments(path)
    _inject_textbox_and_sdt(path)
    _inject_prior_tracked_changes(path)
    _inject_hyperlink(path)
    _inject_ole_object(path)
    _inject_mail_merge(path)


def _inject_footnote_endnote(path: str) -> None:
    content_types = _read_zip_entry(path, "[Content_Types].xml")
    doc_rels = _read_zip_entry(path, "word/_rels/document.xml.rels")

    footnotes_xml = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<w:footnotes xmlns:w="' + W_NS.encode() + b'">'
        b'<w:footnote w:type="separator" w:id="-1"><w:p/></w:footnote>'
        b'<w:footnote w:type="continuationSeparator" w:id="0"><w:p/></w:footnote>'
        b'<w:footnote w:id="1"><w:p><w:r><w:t>Footnote text '
        + MARKERS["footnote"].encode() + b'</w:t></w:r></w:p></w:footnote>'
        b'</w:footnotes>'
    )
    content_types = _add_override_content_type(
        content_types, "/word/footnotes.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
    )
    doc_rels = _add_relationship(
        doc_rels, "rIdFootnotes",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes",
        "footnotes.xml",
    )

    endnotes_xml = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<w:endnotes xmlns:w="' + W_NS.encode() + b'">'
        b'<w:endnote w:type="separator" w:id="-1"><w:p/></w:endnote>'
        b'<w:endnote w:type="continuationSeparator" w:id="0"><w:p/></w:endnote>'
        b'<w:endnote w:id="1"><w:p><w:r><w:t>Endnote text '
        + MARKERS["endnote"].encode() + b'</w:t></w:r></w:p></w:endnote>'
        b'</w:endnotes>'
    )
    content_types = _add_override_content_type(
        content_types, "/word/endnotes.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml",
    )
    doc_rels = _add_relationship(
        doc_rels, "rIdEndnotes",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/endnotes",
        "endnotes.xml",
    )

    _patch_zip(
        path,
        {"[Content_Types].xml": content_types, "word/_rels/document.xml.rels": doc_rels},
        {"word/footnotes.xml": footnotes_xml, "word/endnotes.xml": endnotes_xml},
    )

    document_xml = _read_zip_entry(path, "word/document.xml")
    root = ET.fromstring(document_xml)
    body = root.find(_qn(W_NS, "body"))

    para_refs = ET.SubElement(body, _qn(W_NS, "p"))
    run_fn = ET.SubElement(para_refs, _qn(W_NS, "r"))
    ET.SubElement(run_fn, _qn(W_NS, "footnoteReference"), {_qn(W_NS, "id"): "1"})
    run_en = ET.SubElement(para_refs, _qn(W_NS, "r"))
    ET.SubElement(run_en, _qn(W_NS, "endnoteReference"), {_qn(W_NS, "id"): "1"})

    _write_document_xml(path, root)


def _inject_comments(path: str) -> None:
    """Three comments in the three states `_comment_visibility_map()`
    distinguishes: visible (plain anchored text), hidden (anchored text is
    inside a `w:vanish` run), and deleted (the comment id has no matching
    `w:commentRangeStart` at all -- the "orphaned comment" edge case the
    design doc calls out, which the map classifies as "deleted")."""
    content_types = _read_zip_entry(path, "[Content_Types].xml")
    doc_rels = _read_zip_entry(path, "word/_rels/document.xml.rels")

    comments_xml = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<w:comments xmlns:w="' + W_NS.encode() + b'">'
        b'<w:comment w:id="1" w:author="Tester" w:date="' + PRIOR_REVISION_DATE.encode() + b'">'
        b'<w:p><w:r><w:t>Visible comment body text.</w:t></w:r></w:p></w:comment>'
        b'<w:comment w:id="2" w:author="Tester" w:date="' + PRIOR_REVISION_DATE.encode() + b'">'
        b'<w:p><w:r><w:t>Hidden comment body text.</w:t></w:r></w:p></w:comment>'
        b'<w:comment w:id="3" w:author="Tester" w:date="' + PRIOR_REVISION_DATE.encode() + b'">'
        b'<w:p><w:r><w:t>Orphaned comment with no matching range (deleted).</w:t></w:r></w:p>'
        b'</w:comment>'
        b'</w:comments>'
    )
    content_types = _add_override_content_type(
        content_types, "/word/comments.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
    )
    doc_rels = _add_relationship(
        doc_rels, "rIdComments",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments",
        "comments.xml",
    )
    _patch_zip(
        path,
        {"[Content_Types].xml": content_types, "word/_rels/document.xml.rels": doc_rels},
        {"word/comments.xml": comments_xml},
    )

    document_xml = _read_zip_entry(path, "word/document.xml")
    root = ET.fromstring(document_xml)
    body = root.find(_qn(W_NS, "body"))

    # Comment 1: visible anchor around ordinary body text.
    ET.SubElement(body, _qn(W_NS, "commentRangeStart"), {_qn(W_NS, "id"): "1"})
    para_visible = ET.SubElement(body, _qn(W_NS, "p"))
    run_visible = ET.SubElement(para_visible, _qn(W_NS, "r"))
    t_visible = ET.SubElement(run_visible, _qn(W_NS, "t"))
    t_visible.text = f"Anchored text {MARKERS['comment_visible_anchor']}."
    ET.SubElement(body, _qn(W_NS, "commentRangeEnd"), {_qn(W_NS, "id"): "1"})
    para_ref1 = ET.SubElement(body, _qn(W_NS, "p"))
    run_ref1 = ET.SubElement(para_ref1, _qn(W_NS, "r"))
    ET.SubElement(run_ref1, _qn(W_NS, "commentReference"), {_qn(W_NS, "id"): "1"})

    # Comment 2: hidden anchor -- anchored run carries w:vanish.
    ET.SubElement(body, _qn(W_NS, "commentRangeStart"), {_qn(W_NS, "id"): "2"})
    para_hidden = ET.SubElement(body, _qn(W_NS, "p"))
    run_hidden = ET.SubElement(para_hidden, _qn(W_NS, "r"))
    rpr_hidden = ET.SubElement(run_hidden, _qn(W_NS, "rPr"))
    ET.SubElement(rpr_hidden, _qn(W_NS, "vanish"))
    t_hidden = ET.SubElement(run_hidden, _qn(W_NS, "t"))
    t_hidden.text = f"Anchored text {MARKERS['comment_hidden_anchor']}."
    ET.SubElement(body, _qn(W_NS, "commentRangeEnd"), {_qn(W_NS, "id"): "2"})
    para_ref2 = ET.SubElement(body, _qn(W_NS, "p"))
    run_ref2 = ET.SubElement(para_ref2, _qn(W_NS, "r"))
    ET.SubElement(run_ref2, _qn(W_NS, "commentReference"), {_qn(W_NS, "id"): "2"})

    # Comment 3: intentionally no commentRangeStart/End at all (id "3" is
    # only referenced, never anchored) -- the orphaned-comment "deleted"
    # classification path in `_comment_visibility_map()`.
    para_ref3 = ET.SubElement(body, _qn(W_NS, "p"))
    run_ref3 = ET.SubElement(para_ref3, _qn(W_NS, "r"))
    ET.SubElement(run_ref3, _qn(W_NS, "commentReference"), {_qn(W_NS, "id"): "3"})

    _write_document_xml(path, root)


def _inject_textbox_and_sdt(path: str) -> None:
    document_xml = _read_zip_entry(path, "word/document.xml")
    root = ET.fromstring(document_xml)
    body = root.find(_qn(W_NS, "body"))

    # Legacy VML textbox: w:pict > v:shape > v:textbox > w:txbxContent > w:p
    para_tb = ET.SubElement(body, _qn(W_NS, "p"))
    run_tb = ET.SubElement(para_tb, _qn(W_NS, "r"))
    pict = ET.SubElement(run_tb, _qn(W_NS, "pict"))
    shape = ET.SubElement(pict, _qn(V_NS, "shape"))
    textbox = ET.SubElement(shape, _qn(V_NS, "textbox"))
    txbx_content = ET.SubElement(textbox, _qn(W_NS, "txbxContent"))
    tb_p = ET.SubElement(txbx_content, _qn(W_NS, "p"))
    tb_r = ET.SubElement(tb_p, _qn(W_NS, "r"))
    tb_t = ET.SubElement(tb_r, _qn(W_NS, "t"))
    tb_t.text = f"Textbox text {MARKERS['textbox']}"

    # Block-level content control: w:sdt > w:sdtContent > w:p > w:r > w:t
    sdt = ET.SubElement(body, _qn(W_NS, "sdt"))
    ET.SubElement(sdt, _qn(W_NS, "sdtPr"))
    sdt_content = ET.SubElement(sdt, _qn(W_NS, "sdtContent"))
    sdt_p = ET.SubElement(sdt_content, _qn(W_NS, "p"))
    sdt_r = ET.SubElement(sdt_p, _qn(W_NS, "r"))
    sdt_t = ET.SubElement(sdt_r, _qn(W_NS, "t"))
    sdt_t.text = f"SDT text {MARKERS['sdt_block']}"

    # Inline content control: w:p > w:sdt > w:sdtContent > w:r > w:t
    para_inline_sdt = ET.SubElement(body, _qn(W_NS, "p"))
    inline_sdt = ET.SubElement(para_inline_sdt, _qn(W_NS, "sdt"))
    ET.SubElement(inline_sdt, _qn(W_NS, "sdtPr"))
    inline_sdt_content = ET.SubElement(inline_sdt, _qn(W_NS, "sdtContent"))
    inline_sdt_r = ET.SubElement(inline_sdt_content, _qn(W_NS, "r"))
    inline_sdt_t = ET.SubElement(inline_sdt_r, _qn(W_NS, "t"))
    inline_sdt_t.text = f"Inline SDT text {MARKERS['sdt_inline']}"

    _write_document_xml(path, root)


def _inject_prior_tracked_changes(path: str) -> None:
    """Tracked-changes already present on load (as opposed to ones
    DocxMap.apply_replacements() writes itself). Fixed author/date so the
    fixture stays byte-deterministic."""
    document_xml = _read_zip_entry(path, "word/document.xml")
    root = ET.fromstring(document_xml)
    body = root.find(_qn(W_NS, "body"))

    para_ins = ET.SubElement(body, _qn(W_NS, "p"))
    ins = ET.SubElement(para_ins, _qn(W_NS, "ins"), {
        _qn(W_NS, "id"): "900", _qn(W_NS, "author"): "PriorAuthor",
        _qn(W_NS, "date"): PRIOR_REVISION_DATE,
    })
    run_ins = ET.SubElement(ins, _qn(W_NS, "r"))
    t_ins = ET.SubElement(run_ins, _qn(W_NS, "t"))
    t_ins.text = f"Pre-existing insertion {MARKERS['prior_insert']}."

    para_del = ET.SubElement(body, _qn(W_NS, "p"))
    delel = ET.SubElement(para_del, _qn(W_NS, "del"), {
        _qn(W_NS, "id"): "901", _qn(W_NS, "author"): "PriorAuthor",
        _qn(W_NS, "date"): PRIOR_REVISION_DATE,
    })
    run_del = ET.SubElement(delel, _qn(W_NS, "r"))
    t_del = ET.SubElement(run_del, _qn(W_NS, "delText"))
    t_del.text = "Pre-existing deletion PRIORDELETEMARK_o15."

    _write_document_xml(path, root)


def _inject_hyperlink(path: str) -> None:
    doc_rels = _read_zip_entry(path, "word/_rels/document.xml.rels")
    doc_rels = _add_relationship(
        doc_rels, "rIdHyperlink",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        "https://example.com/redact-me", target_mode="External",
    )
    _patch_zip(path, {"word/_rels/document.xml.rels": doc_rels})

    document_xml = _read_zip_entry(path, "word/document.xml")
    root = ET.fromstring(document_xml)
    body = root.find(_qn(W_NS, "body"))

    para = ET.SubElement(body, _qn(W_NS, "p"))
    hyperlink = ET.SubElement(para, _qn(W_NS, "hyperlink"), {_qn(R_NS, "id"): "rIdHyperlink"})
    run = ET.SubElement(hyperlink, _qn(W_NS, "r"))
    t = ET.SubElement(run, _qn(W_NS, "t"))
    t.text = f"Hyperlink text {MARKERS['hyperlink']}"

    _write_document_xml(path, root)


def _inject_ole_object(path: str) -> None:
    content_types = _read_zip_entry(path, "[Content_Types].xml")
    content_types = _add_default_content_type(
        content_types, "bin", "application/vnd.openxmlformats-officedocument.oleObject",
    )
    doc_rels = _read_zip_entry(path, "word/_rels/document.xml.rels")
    doc_rels = _add_relationship(
        doc_rels, "rIdOle",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject",
        "embeddings/oleObject1.bin",
    )
    _patch_zip(
        path,
        {"[Content_Types].xml": content_types, "word/_rels/document.xml.rels": doc_rels},
        {"word/embeddings/oleObject1.bin": b"FAKE-OLE-BINARY-PAYLOAD-FOR-GOLDEN-FIXTURE"},
    )

    document_xml = _read_zip_entry(path, "word/document.xml")
    root = ET.fromstring(document_xml)
    body = root.find(_qn(W_NS, "body"))

    para = ET.SubElement(body, _qn(W_NS, "p"))
    run = ET.SubElement(para, _qn(W_NS, "r"))
    obj = ET.SubElement(run, _qn(W_NS, "object"))
    ET.SubElement(obj, _qn(O_NS, "OLEObject"), {
        "Type": "Embed", "ProgID": "Word.Document.12",
        "ShapeID": "_x0000_i1025", "DrawAspect": "Content",
        _qn(R_NS, "id"): "rIdOle", "ObjectID": "_1234567890",
    })

    _write_document_xml(path, root)


def _inject_mail_merge(path: str) -> None:
    document_xml = _read_zip_entry(path, "word/document.xml")
    root = ET.fromstring(document_xml)
    body = root.find(_qn(W_NS, "body"))
    para = ET.SubElement(body, _qn(W_NS, "p"))
    fld = ET.SubElement(para, _qn(W_NS, "fldSimple"))
    fld.set(_qn(W_NS, "instr"), " MERGEFIELD FirstName \\* MERGEFORMAT ")
    run = ET.SubElement(fld, _qn(W_NS, "r"))
    text = ET.SubElement(run, _qn(W_NS, "t"))
    text.text = "FirstName"
    updated_doc = ET.tostring(root, encoding="utf-8", xml_declaration=True)

    settings_xml = _read_zip_entry(path, "word/settings.xml")
    settings_root = ET.fromstring(settings_xml)
    mail_merge = ET.SubElement(settings_root, _qn(W_NS, "mailMerge"))
    ET.SubElement(mail_merge, _qn(W_NS, "dataSource"))
    ET.SubElement(mail_merge, _qn(W_NS, "headerSource"))
    updated_settings = ET.tostring(settings_root, encoding="utf-8", xml_declaration=True)

    _patch_zip(path, {"word/document.xml": updated_doc, "word/settings.xml": updated_settings})


def _write_document_xml(path: str, root: ET.Element) -> None:
    updated_doc = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    _patch_zip(path, {"word/document.xml": updated_doc})
