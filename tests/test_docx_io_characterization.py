"""Characterization tests for `docx_io.py` (issue #71).

docs/design/docx_io_package_split.md, §3.2 items 2-5: these are the deeper
unit tests that must be green against the CURRENT, unsplit `docx_io.py`
before Slice 4 (`scan.py`) or Slice 5 (`hardening.py`/`revisions.py`) may
begin. They isolate exactly the surfaces the design doc flags as
least-tested today:

  2. `harden_document()` -- RSID removal, OLE/ActiveX -> placeholder
     replacement (both the "run's only child" and "run has other content"
     cases), and `scrub_all_images`.
  3. `_build()`/the `_scan_*` family, isolated from `apply_replacements()`
     -- `.text`/`.index` asserted directly for textboxes (`_scan_drawing_tag`)
     and nested tables (`_scan_table_xml`'s recursion).
  4. `_comment_visibility_map()` -- hidden/deleted/visible plus the
     orphaned-comment-id edge case.
  5. The `_metadata_settings` -> `save()` order-dependency contract that
     gates whether ZIP-level hardening (`_rewrite_docx_zip`) runs at all.

This module makes no edits to `docx_io.py` -- tests and fixtures only.
Every assertion is against parsed XML element trees (python-docx/lxml) or
the `.text`/`.index`/dict values `DocxMap` builds in memory, never raw
bytes and never just "no exception raised".

Run just this suite with:

    PYTHONPATH=src/python python3 -m pytest tests/test_docx_io_characterization.py -q
"""
from __future__ import annotations

import os
import tempfile
import zipfile
from unittest import mock
from xml.etree import ElementTree as ET

import pytest

from docx import Document
from docx.enum.section import WD_HEADER_FOOTER, WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from marcut.docx_io import DocxMap, MetadataCleaningSettings

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def _qn(ns: str, tag: str) -> str:
    return f"{{{ns}}}{tag}"


def _new_docmap() -> tuple[Document, DocxMap]:
    """A fresh in-memory Document wrapped by a DocxMap. Building the
    DocxMap directly from the live python-docx object (no save/reload
    round trip) is enough to exercise `_build()`/`_scan_*`/`harden_document`
    in isolation -- `apply_replacements()` is never called here."""
    doc = Document()
    return doc, DocxMap(doc)


def _insert_before_sectpr(body, element) -> None:
    sect_pr = body.find(qn("w:sectPr"))
    if sect_pr is not None:
        sect_pr.addprevious(element)
    else:
        body.append(element)


def _et_insert_before_sectpr(body: ET.Element, element: ET.Element) -> None:
    """`xml.etree` analogue of `_insert_before_sectpr` (ElementTree has no
    `addprevious`) for fixtures that patch a serialized document.xml:
    keeps `w:sectPr` the last child of `w:body`, as WordprocessingML
    requires, when a fixture adds block-level content."""
    sect_pr = body.find(_qn(W_NS, "sectPr"))
    if sect_pr is not None:
        body.insert(list(body).index(sect_pr), element)
    else:
        body.append(element)


# ---------------------------------------------------------------------------
# 2. harden_document()
# ---------------------------------------------------------------------------


class TestHardenDocumentRSIDs:
    def test_rsid_attributes_removed_from_paragraph_and_run(self):
        doc, dm = _new_docmap()
        p = doc.add_paragraph()
        p._p.set(qn("w:rsidR"), "00AB12CD")
        p._p.set(qn("w:rsidRDefault"), "00AB12CD")
        r = OxmlElement("w:r")
        r.set(qn("w:rsidR"), "00AB12CD")
        t = OxmlElement("w:t")
        t.text = "hello"
        r.append(t)
        p._p.append(r)

        dm.harden_document(settings=MetadataCleaningSettings())

        for elem in p._p.iter():
            for key in elem.attrib:
                assert "rsid" not in key.lower(), f"leftover rsid attr {key!r} on {elem.tag}"

    def test_rsid_attributes_kept_when_clean_rsids_disabled(self):
        doc, dm = _new_docmap()
        p = doc.add_paragraph()
        p._p.set(qn("w:rsidR"), "00AB12CD")

        dm.harden_document(settings=MetadataCleaningSettings(clean_rsids=False))

        assert p._p.get(qn("w:rsidR")) == "00AB12CD"


class TestHardenDocumentOleActiveX:
    """Pins the two branches of the OLE/ActiveX replacement logic in
    `harden_document()`'s nested `_harden_element()`: when the object is
    literally the run's only child (`len(parent) == 0` after removal), a
    placeholder `<w:t>` is appended into that same run; when the run has
    other content (e.g. `w:rPr`), the object is removed but -- as the
    code stands today -- NO placeholder is added at all. That second
    case is a real, currently-shipping quirk (silent content loss), not a
    hypothesis; it must be preserved exactly as-is by any refactor unless
    a separate ticket deliberately changes it."""

    def test_object_as_runs_only_child_gets_placeholder_in_same_run(self):
        doc, dm = _new_docmap()
        p = doc.add_paragraph()
        r = OxmlElement("w:r")
        r.append(OxmlElement("w:object"))
        p._p.append(r)

        dm.harden_document(settings=MetadataCleaningSettings())

        runs = p._p.findall(qn("w:r"))
        assert len(runs) == 1
        assert runs[0].find(qn("w:object")) is None
        t = runs[0].find(qn("w:t"))
        assert t is not None
        assert t.text == "[REDACTED DATA OBJECT]"

    def test_object_in_run_with_other_content_is_removed_without_placeholder(self):
        doc, dm = _new_docmap()
        p = doc.add_paragraph()
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        rpr.append(OxmlElement("w:b"))
        r.append(rpr)
        r.append(OxmlElement("w:object"))
        p._p.append(r)

        dm.harden_document(settings=MetadataCleaningSettings())

        runs = p._p.findall(qn("w:r"))
        assert len(runs) == 1
        assert runs[0].find(qn("w:object")) is None
        # Characterizes the current behavior exactly: no placeholder text
        # is inserted in this branch, the run just loses its object.
        assert runs[0].find(qn("w:t")) is None
        assert runs[0].find(qn("w:rPr")) is not None

    def test_object_not_wrapped_in_a_run_is_replaced_by_new_placeholder_run_in_place(self):
        doc, dm = _new_docmap()
        p = doc.add_paragraph()
        obj = OxmlElement("w:object")
        p._p.append(obj)
        after_run = OxmlElement("w:r")
        after_t = OxmlElement("w:t")
        after_t.text = "AFTER"
        after_run.append(after_t)
        p._p.append(after_run)

        dm.harden_document(settings=MetadataCleaningSettings())

        children = [c for c in p._p if c.tag == qn("w:r")]
        assert len(children) == 2
        assert children[0].find(qn("w:t")).text == "[REDACTED DATA OBJECT]"
        assert children[1].find(qn("w:t")).text == "AFTER"

    def test_control_tag_replaced_when_only_ole_objects_enabled(self):
        """`w:control` (ActiveX) is gated by `clean_ole_objects OR
        clean_activex` -- either flag alone is sufficient."""
        doc, dm = _new_docmap()
        p = doc.add_paragraph()
        r = OxmlElement("w:r")
        r.append(OxmlElement("w:control"))
        p._p.append(r)

        dm.harden_document(
            settings=MetadataCleaningSettings(clean_ole_objects=True, clean_activex=False)
        )

        runs = p._p.findall(qn("w:r"))
        assert runs[0].find(qn("w:control")) is None
        assert runs[0].find(qn("w:t")).text == "[REDACTED DATA OBJECT]"

    def test_object_and_control_untouched_when_both_disabled(self):
        doc, dm = _new_docmap()
        p = doc.add_paragraph()
        r = OxmlElement("w:r")
        r.append(OxmlElement("w:object"))
        p._p.append(r)

        dm.harden_document(
            settings=MetadataCleaningSettings(clean_ole_objects=False, clean_activex=False)
        )

        runs = p._p.findall(qn("w:r"))
        assert runs[0].find(qn("w:object")) is not None


class TestHardenDocumentScrubAllImages:
    def test_scrub_all_images_true_removes_drawing_and_pict(self):
        doc, dm = _new_docmap()
        p = doc.add_paragraph()
        r = OxmlElement("w:r")
        r.append(OxmlElement("w:drawing"))
        r.append(OxmlElement("w:pict"))
        p._p.append(r)

        dm.harden_document(scrub_all_images=True, settings=MetadataCleaningSettings())

        runs = p._p.findall(qn("w:r"))
        assert runs[0].find(qn("w:drawing")) is None
        assert runs[0].find(qn("w:pict")) is None

    def test_scrub_all_images_false_leaves_drawing_untouched(self):
        doc, dm = _new_docmap()
        p = doc.add_paragraph()
        r = OxmlElement("w:r")
        r.append(OxmlElement("w:drawing"))
        p._p.append(r)

        dm.harden_document(scrub_all_images=False, settings=MetadataCleaningSettings())

        runs = p._p.findall(qn("w:r"))
        assert runs[0].find(qn("w:drawing")) is not None


# ---------------------------------------------------------------------------
# 3. _build() / _scan_* family
# ---------------------------------------------------------------------------


class TestScanIndexing:
    def test_scan_textbox_via_scan_drawing_tag(self):
        doc = Document()
        p = doc.add_paragraph()
        drawing = OxmlElement("w:drawing")
        txbx = OxmlElement("w:txbxContent")
        tb_p = OxmlElement("w:p")
        tb_r = OxmlElement("w:r")
        tb_t = OxmlElement("w:t")
        tb_t.text = "textbox body"
        tb_r.append(tb_t)
        tb_p.append(tb_r)
        txbx.append(tb_p)
        drawing.append(txbx)
        p._p.append(drawing)

        dm = DocxMap(doc)

        assert "textbox body" in dm.text
        # Reconstructing from .index must reproduce .text exactly -- this
        # is the invariant a scan.py/hardening.py split must preserve.
        assert _reconstruct_text_from_index(dm) == dm.text

    def test_scan_textbox_nested_in_run_drawing_via_scan_run_contents(self):
        """The path Word actually emits: `w:p` -> `w:r` -> `w:drawing` ->
        `w:txbxContent`, which is reached through `_scan_run_contents`
        (not the `w:drawing`-as-direct-child-of-`w:p` branch the sibling
        test above covers). This is the path a move of the iteration
        helpers into scan.py would break if `_scan_run_contents` weren't
        carried along with `_scan_paragraph`."""
        doc = Document()
        p = doc.add_paragraph()
        outer_run = OxmlElement("w:r")
        drawing = OxmlElement("w:drawing")
        txbx = OxmlElement("w:txbxContent")
        tb_p = OxmlElement("w:p")
        tb_r = OxmlElement("w:r")
        tb_t = OxmlElement("w:t")
        tb_t.text = "run-drawing textbox body"
        tb_r.append(tb_t)
        tb_p.append(tb_r)
        txbx.append(tb_p)
        drawing.append(txbx)
        outer_run.append(drawing)
        p._p.append(outer_run)

        dm = DocxMap(doc)

        assert "run-drawing textbox body" in dm.text
        assert _reconstruct_text_from_index(dm) == dm.text

    def test_scan_textbox_nested_in_run_pict_via_scan_run_contents(self):
        """Same path as above but via the legacy VML `w:pict` wrapper
        instead of `w:drawing` -- `_scan_run_contents` iterates both
        tags identically."""
        doc = Document()
        p = doc.add_paragraph()
        outer_run = OxmlElement("w:r")
        pict = OxmlElement("w:pict")
        txbx = OxmlElement("w:txbxContent")
        tb_p = OxmlElement("w:p")
        tb_r = OxmlElement("w:r")
        tb_t = OxmlElement("w:t")
        tb_t.text = "run-pict textbox body"
        tb_r.append(tb_t)
        tb_p.append(tb_r)
        txbx.append(tb_p)
        pict.append(txbx)
        outer_run.append(pict)
        p._p.append(outer_run)

        dm = DocxMap(doc)

        assert "run-pict textbox body" in dm.text
        assert _reconstruct_text_from_index(dm) == dm.text

    def test_scan_block_level_content_control_sdt(self):
        doc = Document()
        body = doc.element.body
        sdt = OxmlElement("w:sdt")
        sdt.append(OxmlElement("w:sdtPr"))
        sdt_content = OxmlElement("w:sdtContent")
        sdt_p = OxmlElement("w:p")
        sdt_r = OxmlElement("w:r")
        sdt_t = OxmlElement("w:t")
        sdt_t.text = "block sdt body"
        sdt_r.append(sdt_t)
        sdt_p.append(sdt_r)
        sdt_content.append(sdt_p)
        sdt.append(sdt_content)
        _insert_before_sectpr(body, sdt)

        dm = DocxMap(doc)

        assert "block sdt body" in dm.text
        assert _reconstruct_text_from_index(dm) == dm.text

    def test_scan_inline_content_control_sdt(self):
        doc = Document()
        p = doc.add_paragraph()
        sdt = OxmlElement("w:sdt")
        sdt.append(OxmlElement("w:sdtPr"))
        sdt_content = OxmlElement("w:sdtContent")
        sdt_r = OxmlElement("w:r")
        sdt_t = OxmlElement("w:t")
        sdt_t.text = "inline sdt body"
        sdt_r.append(sdt_t)
        sdt_content.append(sdt_r)
        sdt.append(sdt_content)
        p._p.append(sdt)

        dm = DocxMap(doc)

        assert "inline sdt body" in dm.text
        assert _reconstruct_text_from_index(dm) == dm.text

    def test_scan_table_and_nested_table_via_scan_table_xml(self):
        """Pins `_scan_table_xml`'s recursion exactly as it behaves today.

        `_scan_table_xml` walks `tbl_element.iter(qn('w:tr'))` /
        `tr.iter(qn('w:tc'))` -- both are RECURSIVE lxml `.iter()` calls,
        so for an outer table containing a nested table, the outer call
        yields both the outer `<w:tr>` and the nested table's `<w:tr>` as
        separate top-level iterations, on top of the explicit recursive
        call the code also makes when it sees a `w:tbl` child. The net
        effect is that today's code scans the nested cell's paragraph
        THREE times, not once. This is exactly the kind of fragile,
        surprising path the design doc calls out (docs/design/
        docx_io_package_split.md §3.2 item 3) -- characterizing it here
        means a `scan.py` extraction can't silently "fix" (i.e. change)
        it without this test failing loudly.
        """
        doc = Document()
        outer = doc.add_table(rows=1, cols=1)
        outer_cell = outer.cell(0, 0)
        nested = outer_cell.add_table(rows=1, cols=1)
        nested.cell(0, 0).text = "nested-text"

        dm = DocxMap(doc)

        assert dm.text.count("nested-text") == 3
        assert dm.text == "\nnested-text\n\nnested-text\nnested-text\n"
        assert _reconstruct_text_from_index(dm) == dm.text

    def test_scan_simple_table_single_cell_no_nesting(self):
        """Baseline sanity check: a table with no nested table is scanned
        exactly once (the triple-count above is specific to nesting)."""
        doc = Document()
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "plain-cell-text"

        dm = DocxMap(doc)

        assert dm.text.count("plain-cell-text") == 1
        assert _reconstruct_text_from_index(dm) == dm.text


def _build_footnote_endnote_fixture(path: str) -> None:
    """A minimal document with a footnote and an endnote wired up as real
    OPC parts/relationships (docx has no first-class API for these), so
    `DocxMap.load()` exercises the same generic, non-specialized `Part`
    path (no `.element` attribute) `_build()` falls back to -- the exact
    path that appends to `self.detached_parts`."""
    doc = Document()
    doc.add_paragraph("placeholder body")
    doc.save(path)

    content_types = _read_zip_entry(path, "[Content_Types].xml")
    doc_rels = _read_zip_entry(path, "word/_rels/document.xml.rels")

    footnotes_xml = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<w:footnotes xmlns:w="' + W_NS.encode() + b'">'
        b'<w:footnote w:id="1"><w:p><w:r><w:t>footnote marker text</w:t></w:r></w:p></w:footnote>'
        b'</w:footnotes>'
    )
    content_types = _add_content_type_override(
        content_types, "/word/footnotes.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
    )
    doc_rels = _add_relationship(
        doc_rels, "rIdFootnotes",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes",
        "footnotes.xml",
    )

    endnotes_xml = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<w:endnotes xmlns:w="' + W_NS.encode() + b'">'
        b'<w:endnote w:id="1"><w:p><w:r><w:t>endnote marker text</w:t></w:r></w:p></w:endnote>'
        b'</w:endnotes>'
    )
    content_types = _add_content_type_override(
        content_types, "/word/endnotes.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml",
    )
    doc_rels = _add_relationship(
        doc_rels, "rIdEndnotes",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/endnotes",
        "endnotes.xml",
    )

    _patch_zip(
        path,
        {"[Content_Types].xml": content_types, "word/_rels/document.xml.rels": doc_rels},
        {"word/footnotes.xml": footnotes_xml, "word/endnotes.xml": endnotes_xml},
    )

    document_xml = _read_zip_entry(path, "word/document.xml")
    root = ET.fromstring(document_xml)
    body = root.find(_qn(W_NS, "body"))
    p = ET.Element(_qn(W_NS, "p"))
    r_fn = ET.SubElement(p, _qn(W_NS, "r"))
    ET.SubElement(r_fn, _qn(W_NS, "footnoteReference"), {_qn(W_NS, "id"): "1"})
    r_en = ET.SubElement(p, _qn(W_NS, "r"))
    ET.SubElement(r_en, _qn(W_NS, "endnoteReference"), {_qn(W_NS, "id"): "1"})
    _et_insert_before_sectpr(body, p)
    _patch_zip(path, {"word/document.xml": ET.tostring(root, encoding="utf-8", xml_declaration=True)})


class TestBuildPerContainerType:
    """`_build()`'s per-container-type behavior, isolated from the
    end-to-end golden harness (#70): the two distinct header/footer skip
    paths (the `is_linked_to_previous` check, and the `scanned_elements`
    id()-based dedup that only fires when two UNLINKED sections resolve
    to the same header/footer part), and the footnote/endnote
    generic-part path that appends to `self.detached_parts` -- the
    shared-state hazard the design doc names for Slice 4."""

    def test_headers_and_footers_indexed_once_each_across_sections(self):
        """The `is_linked_to_previous` skip. A second section left
        linked-to-previous (python-docx's default after `add_section`)
        has no `w:headerReference`/`w:footerReference` of its own, so
        `_build()` skips it on the linked flag before the
        `scanned_elements` id() set is ever consulted -- it never touches
        `header._element`, which for a linked header would resolve to the
        prior section's part. The id()-dedup branch is exercised
        separately by the shared-`r:id` test below."""
        doc = Document()
        doc.add_paragraph("body text")
        sec1 = doc.sections[0]
        sec1.header.is_linked_to_previous = False
        sec1.header.paragraphs[0].text = "header marker text"
        sec1.footer.is_linked_to_previous = False
        sec1.footer.paragraphs[0].text = "footer marker text"

        # A second section left linked-to-previous (python-docx's
        # default): no header/footer reference of its own, so `_build()`
        # skips it on the linked flag and never reaches the id() dedup.
        doc.add_section(WD_SECTION.NEW_PAGE)
        doc.add_paragraph("more body text")

        dm = DocxMap(doc)

        assert dm.text.count("header marker text") == 1
        assert dm.text.count("footer marker text") == 1
        assert _reconstruct_text_from_index(dm) == dm.text

    def test_unlinked_sections_sharing_one_header_part_deduped_by_element_id(self):
        """The `scanned_elements` id()-dedup branch. Two UNLINKED sections
        whose `w:headerReference` elements carry the same `r:id` both
        resolve (via `DocumentPart.header_part(rId)`) to the one
        `HeaderPart`, hence the same `w:hdr` element object. Neither is
        skipped on `is_linked_to_previous`, so only the id() set stops
        the second scan; without it the header text is indexed twice."""
        doc = Document()
        doc.add_paragraph("body text")
        first = doc.sections[0]
        first.header.is_linked_to_previous = False
        first.header.paragraphs[0].text = "shared header text"
        shared_rid = first._sectPr.get_headerReference(WD_HEADER_FOOTER.PRIMARY).rId

        # `add_section` clones the sentinel sectPr (header reference
        # included) into a paragraph as section 0 and strips the
        # references off the sentinel, which becomes section 1. Point
        # section 1 at section 0's header part explicitly.
        doc.add_section(WD_SECTION.NEW_PAGE)
        doc.add_paragraph("more body text")
        sec1, sec2 = doc.sections[0], doc.sections[1]
        sec2._sectPr.add_headerReference(WD_HEADER_FOOTER.PRIMARY, shared_rid)
        assert not sec1.header.is_linked_to_previous
        assert not sec2.header.is_linked_to_previous
        assert sec1.header._element is sec2.header._element

        dm = DocxMap(doc)

        assert dm.text.count("shared header text") == 1
        assert _reconstruct_text_from_index(dm) == dm.text

    def test_footnotes_and_endnotes_indexed_via_generic_detached_part(self, tmp_path):
        path = str(tmp_path / "fn_en.docx")
        _build_footnote_endnote_fixture(path)

        dm = DocxMap.load(path)

        assert "footnote marker text" in dm.text
        assert "endnote marker text" in dm.text
        assert _reconstruct_text_from_index(dm) == dm.text

        detached_names = {str(part.partname) for part, _root in dm.detached_parts}
        assert "/word/footnotes.xml" in detached_names
        assert "/word/endnotes.xml" in detached_names


def _reconstruct_text_from_index(dm: DocxMap) -> str:
    """Walks `dm.index` and rebuilds the string it should have produced,
    entry by entry, so `.text`/`.index` consistency is asserted directly
    rather than assumed."""
    chars = []
    for entry in dm.index:
        para, run, i = entry
        if para == "break":
            chars.append("\n")
        else:
            chars.append((run.text or "")[i])
    return "".join(chars)


# ---------------------------------------------------------------------------
# 4. _comment_visibility_map()
# ---------------------------------------------------------------------------


def _read_zip_entry(path: str, name: str) -> bytes:
    with zipfile.ZipFile(path) as zf:
        return zf.read(name)


def _patch_zip(path: str, updates: dict, new_entries: dict | None = None) -> None:
    new_entries = new_entries or {}
    tmp = path + ".tmp"
    with zipfile.ZipFile(path, "r") as zin:
        existing = {i.filename for i in zin.infolist()}
        with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = updates.get(item.filename, zin.read(item.filename))
                zout.writestr(item, data)
            for name, data in new_entries.items():
                if name not in existing:
                    zout.writestr(name, data)
    os.replace(tmp, path)


def _add_content_type_override(content_types_xml: bytes, part_name: str, content_type: str) -> bytes:
    root = ET.fromstring(content_types_xml)
    for override in root.findall(_qn(CT_NS, "Override")):
        if override.get("PartName") == part_name:
            return content_types_xml
    ET.SubElement(root, _qn(CT_NS, "Override"), {"PartName": part_name, "ContentType": content_type})
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _add_relationship(rels_xml: bytes, rel_id: str, rel_type: str, target: str) -> bytes:
    root = ET.fromstring(rels_xml)
    for rel in root.findall(_qn(REL_NS, "Relationship")):
        if rel.get("Id") == rel_id:
            return rels_xml
    ET.SubElement(root, _qn(REL_NS, "Relationship"), {"Id": rel_id, "Type": rel_type, "Target": target})
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _build_comment_visibility_fixture(path: str) -> None:
    """One document exercising all three `_comment_visibility_map()`
    states plus the id-mismatch edge cases in both directions:
      id 1 -- visible (plain anchored text)
      id 2 -- hidden (anchored run carries w:vanish)
      id 3 -- deleted (anchored range sits inside a w:del)
      id 4 -- inverse orphan: declared in comments.xml but no
              commentRangeStart/End anywhere in the document at all
              (the `not anchor_info.get("has_start")` branch)
      id 5 -- the ticket's actual named edge case: a commentRangeStart/
              End (+ commentReference) in the document with NO matching
              `w:comment` entry in comments.xml at all. Because
              `_comment_visibility_map()` only iterates `comment_ids`
              (the ids declared in comments.xml), id 5 never becomes a
              key of the returned map -- it is silently absent, not
              classified "deleted".
      id 6 -- deleted via `w:moveFrom` ancestry on the anchored text
              (the `_has_ancestor(elem, qn('w:moveFrom'))` branch)
      id 7 -- hidden via `w:webHidden` (as opposed to `w:vanish`)
      id 8 -- deleted because the `commentRangeStart` element itself
              sits inside a `w:del` (the `_has_ancestor(elem, qn('w:del'))`
              check on the start tag, not on the anchored text)
    """
    doc = Document()
    doc.add_paragraph("placeholder body")
    doc.save(path)

    content_types = _read_zip_entry(path, "[Content_Types].xml")
    doc_rels = _read_zip_entry(path, "word/_rels/document.xml.rels")

    comments_xml = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<w:comments xmlns:w="' + W_NS.encode() + b'">'
        b'<w:comment w:id="1" w:author="T" w:date="2020-01-01T00:00:00Z">'
        b'<w:p><w:r><w:t>visible comment</w:t></w:r></w:p></w:comment>'
        b'<w:comment w:id="2" w:author="T" w:date="2020-01-01T00:00:00Z">'
        b'<w:p><w:r><w:t>hidden comment</w:t></w:r></w:p></w:comment>'
        b'<w:comment w:id="3" w:author="T" w:date="2020-01-01T00:00:00Z">'
        b'<w:p><w:r><w:t>deleted comment</w:t></w:r></w:p></w:comment>'
        b'<w:comment w:id="4" w:author="T" w:date="2020-01-01T00:00:00Z">'
        b'<w:p><w:r><w:t>orphaned comment</w:t></w:r></w:p></w:comment>'
        b'<w:comment w:id="6" w:author="T" w:date="2020-01-01T00:00:00Z">'
        b'<w:p><w:r><w:t>moveFrom comment</w:t></w:r></w:p></w:comment>'
        b'<w:comment w:id="7" w:author="T" w:date="2020-01-01T00:00:00Z">'
        b'<w:p><w:r><w:t>webHidden comment</w:t></w:r></w:p></w:comment>'
        b'<w:comment w:id="8" w:author="T" w:date="2020-01-01T00:00:00Z">'
        b'<w:p><w:r><w:t>start-in-del comment</w:t></w:r></w:p></w:comment>'
        b'</w:comments>'
    )
    content_types = _add_content_type_override(
        content_types, "/word/comments.xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
    )
    doc_rels = _add_relationship(
        doc_rels, "rIdComments",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments",
        "comments.xml",
    )
    _patch_zip(
        path,
        {"[Content_Types].xml": content_types, "word/_rels/document.xml.rels": doc_rels},
        {"word/comments.xml": comments_xml},
    )

    document_xml = _read_zip_entry(path, "word/document.xml")
    root = ET.fromstring(document_xml)
    body = root.find(_qn(W_NS, "body"))

    # Each comment is built the way Word serializes it: the range markers
    # and the trailing commentReference run live INSIDE the anchored
    # paragraph, and every paragraph is inserted before w:sectPr (which
    # must remain the last child of w:body).

    # id 1: visible.
    p1 = ET.Element(_qn(W_NS, "p"))
    ET.SubElement(p1, _qn(W_NS, "commentRangeStart"), {_qn(W_NS, "id"): "1"})
    r1 = ET.SubElement(p1, _qn(W_NS, "r"))
    t1 = ET.SubElement(r1, _qn(W_NS, "t"))
    t1.text = "visible anchor text"
    ET.SubElement(p1, _qn(W_NS, "commentRangeEnd"), {_qn(W_NS, "id"): "1"})
    ref1_r = ET.SubElement(p1, _qn(W_NS, "r"))
    ET.SubElement(ref1_r, _qn(W_NS, "commentReference"), {_qn(W_NS, "id"): "1"})
    _et_insert_before_sectpr(body, p1)

    # id 2: hidden -- anchored run carries w:vanish.
    p2 = ET.Element(_qn(W_NS, "p"))
    ET.SubElement(p2, _qn(W_NS, "commentRangeStart"), {_qn(W_NS, "id"): "2"})
    r2 = ET.SubElement(p2, _qn(W_NS, "r"))
    rpr2 = ET.SubElement(r2, _qn(W_NS, "rPr"))
    ET.SubElement(rpr2, _qn(W_NS, "vanish"))
    t2 = ET.SubElement(r2, _qn(W_NS, "t"))
    t2.text = "hidden anchor text"
    ET.SubElement(p2, _qn(W_NS, "commentRangeEnd"), {_qn(W_NS, "id"): "2"})
    ref2_r = ET.SubElement(p2, _qn(W_NS, "r"))
    ET.SubElement(ref2_r, _qn(W_NS, "commentReference"), {_qn(W_NS, "id"): "2"})
    _et_insert_before_sectpr(body, p2)

    # id 3: deleted -- anchored range sits inside a w:del.
    p3 = ET.Element(_qn(W_NS, "p"))
    ET.SubElement(p3, _qn(W_NS, "commentRangeStart"), {_qn(W_NS, "id"): "3"})
    del3 = ET.SubElement(p3, _qn(W_NS, "del"), {
        _qn(W_NS, "id"): "500", _qn(W_NS, "author"): "T", _qn(W_NS, "date"): "2020-01-01T00:00:00Z",
    })
    r3 = ET.SubElement(del3, _qn(W_NS, "r"))
    t3 = ET.SubElement(r3, _qn(W_NS, "delText"))
    t3.text = "deleted anchor text"
    ET.SubElement(p3, _qn(W_NS, "commentRangeEnd"), {_qn(W_NS, "id"): "3"})
    ref3_r = ET.SubElement(p3, _qn(W_NS, "r"))
    ET.SubElement(ref3_r, _qn(W_NS, "commentReference"), {_qn(W_NS, "id"): "3"})
    _et_insert_before_sectpr(body, p3)

    # id 4: inverse orphan -- declared in comments.xml, no start/end
    # anywhere, only a dangling commentReference. This is the INVERSE of
    # the ticket's named edge case (see id 5 below): here the id exists
    # in comments.xml but not in the document, so it still becomes a key
    # of the map (classified "deleted" via `not has_start`).
    p4 = ET.Element(_qn(W_NS, "p"))
    ref4_r = ET.SubElement(p4, _qn(W_NS, "r"))
    ET.SubElement(ref4_r, _qn(W_NS, "commentReference"), {_qn(W_NS, "id"): "4"})
    _et_insert_before_sectpr(body, p4)

    # id 5: the ticket's actual named edge case -- a full range/reference
    # in the document, but NO matching `w:comment` entry in comments.xml.
    # Deliberately NOT added to comment_ids below.
    p5 = ET.Element(_qn(W_NS, "p"))
    ET.SubElement(p5, _qn(W_NS, "commentRangeStart"), {_qn(W_NS, "id"): "5"})
    r5 = ET.SubElement(p5, _qn(W_NS, "r"))
    t5 = ET.SubElement(r5, _qn(W_NS, "t"))
    t5.text = "no comments.xml entry anchor text"
    ET.SubElement(p5, _qn(W_NS, "commentRangeEnd"), {_qn(W_NS, "id"): "5"})
    ref5_r = ET.SubElement(p5, _qn(W_NS, "r"))
    ET.SubElement(ref5_r, _qn(W_NS, "commentReference"), {_qn(W_NS, "id"): "5"})
    _et_insert_before_sectpr(body, p5)

    # id 6: deleted via w:moveFrom ancestry on the anchored text (as
    # opposed to id 3's w:del ancestry).
    p6 = ET.Element(_qn(W_NS, "p"))
    ET.SubElement(p6, _qn(W_NS, "commentRangeStart"), {_qn(W_NS, "id"): "6"})
    movefrom6 = ET.SubElement(p6, _qn(W_NS, "moveFrom"), {
        _qn(W_NS, "id"): "600", _qn(W_NS, "author"): "T", _qn(W_NS, "date"): "2020-01-01T00:00:00Z",
    })
    r6 = ET.SubElement(movefrom6, _qn(W_NS, "r"))
    t6 = ET.SubElement(r6, _qn(W_NS, "delText"))
    t6.text = "moved-from anchor text"
    ET.SubElement(p6, _qn(W_NS, "commentRangeEnd"), {_qn(W_NS, "id"): "6"})
    ref6_r = ET.SubElement(p6, _qn(W_NS, "r"))
    ET.SubElement(ref6_r, _qn(W_NS, "commentReference"), {_qn(W_NS, "id"): "6"})
    _et_insert_before_sectpr(body, p6)

    # id 7: hidden via w:webHidden (as opposed to id 2's w:vanish).
    p7 = ET.Element(_qn(W_NS, "p"))
    ET.SubElement(p7, _qn(W_NS, "commentRangeStart"), {_qn(W_NS, "id"): "7"})
    r7 = ET.SubElement(p7, _qn(W_NS, "r"))
    rpr7 = ET.SubElement(r7, _qn(W_NS, "rPr"))
    ET.SubElement(rpr7, _qn(W_NS, "webHidden"))
    t7 = ET.SubElement(r7, _qn(W_NS, "t"))
    t7.text = "webHidden anchor text"
    ET.SubElement(p7, _qn(W_NS, "commentRangeEnd"), {_qn(W_NS, "id"): "7"})
    ref7_r = ET.SubElement(p7, _qn(W_NS, "r"))
    ET.SubElement(ref7_r, _qn(W_NS, "commentReference"), {_qn(W_NS, "id"): "7"})
    _et_insert_before_sectpr(body, p7)

    # id 8: deleted because the commentRangeStart element ITSELF sits
    # inside a w:del -- the anchored text after it is plain, so this
    # isolates the start-tag ancestry check from the anchored-text check
    # id 3/6 already cover.
    p8 = ET.Element(_qn(W_NS, "p"))
    del8 = ET.SubElement(p8, _qn(W_NS, "del"), {
        _qn(W_NS, "id"): "800", _qn(W_NS, "author"): "T", _qn(W_NS, "date"): "2020-01-01T00:00:00Z",
    })
    ET.SubElement(del8, _qn(W_NS, "commentRangeStart"), {_qn(W_NS, "id"): "8"})
    r8 = ET.SubElement(p8, _qn(W_NS, "r"))
    t8 = ET.SubElement(r8, _qn(W_NS, "t"))
    t8.text = "start-in-del anchor text"
    ET.SubElement(p8, _qn(W_NS, "commentRangeEnd"), {_qn(W_NS, "id"): "8"})
    ref8_r = ET.SubElement(p8, _qn(W_NS, "r"))
    ET.SubElement(ref8_r, _qn(W_NS, "commentReference"), {_qn(W_NS, "id"): "8"})
    _et_insert_before_sectpr(body, p8)

    updated_doc = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    _patch_zip(path, {"word/document.xml": updated_doc})


class TestCommentVisibilityMap:
    @pytest.fixture()
    def visibility_map(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "comments.docx")
            _build_comment_visibility_fixture(path)
            dm = DocxMap.load(path)
            yield dm._comment_visibility_map()

    def test_visible_comment_classified_visible(self, visibility_map):
        assert visibility_map["1"] == "visible"

    def test_hidden_comment_classified_hidden(self, visibility_map):
        assert visibility_map["2"] == "hidden"

    def test_deleted_comment_classified_deleted(self, visibility_map):
        assert visibility_map["3"] == "deleted"

    def test_orphaned_comment_id_with_no_matching_range_classified_deleted(self, visibility_map):
        """A comment id present in comments.xml with no matching
        commentRangeStart anywhere in the document is classified
        'deleted' (the `not anchor_info.get("has_start")` branch)."""
        assert visibility_map["4"] == "deleted"

    def test_range_start_with_no_comments_xml_entry_is_absent_from_map(self, visibility_map):
        """The ticket's named edge case: a commentRangeStart/End with an
        id that has no matching `w:comment` entry in comments.xml at
        all. `_comment_visibility_map()` only iterates the ids declared
        in comments.xml, so id 5 is silently absent from the returned
        map -- it is NOT classified 'deleted'."""
        assert "5" not in visibility_map

    def test_movefrom_ancestry_classified_deleted(self, visibility_map):
        assert visibility_map["6"] == "deleted"

    def test_webhidden_run_classified_hidden(self, visibility_map):
        assert visibility_map["7"] == "hidden"

    def test_range_start_inside_del_classified_deleted(self, visibility_map):
        """The commentRangeStart element itself sits inside a w:del
        (the anchored text after it is plain) -- isolates the start-tag
        ancestry check at `_has_ancestor(elem, qn('w:del'))` from the
        anchored-text check id 3/6 already cover."""
        assert visibility_map["8"] == "deleted"

    def test_all_expected_ids_present(self, visibility_map):
        assert set(visibility_map.keys()) == {"1", "2", "3", "4", "6", "7", "8"}


# ---------------------------------------------------------------------------
# 5. Order-dependency: scrub_metadata() -> save() gates _rewrite_docx_zip()
# ---------------------------------------------------------------------------


class TestSaveHardeningOrderDependency:
    def test_hardening_runs_when_scrub_metadata_called_before_save(self, tmp_path):
        doc = Document()
        doc.add_paragraph("body text")
        dm = DocxMap(doc)
        settings = MetadataCleaningSettings()
        dm.scrub_metadata(settings)

        out_path = str(tmp_path / "out.docx")
        with mock.patch.object(DocxMap, "_rewrite_docx_zip") as rewrite:
            dm.save(out_path)
        rewrite.assert_called_once()
        # And the exact settings object threaded through unchanged --
        # `is` here (not `isinstance`) is the point: a fresh default
        # instance would satisfy `isinstance` too but would not prove
        # anything was "threaded through".
        assert rewrite.call_args.args[0] == out_path
        assert rewrite.call_args.args[1] is settings

    def test_hardening_skipped_when_save_called_without_prior_scrub_metadata(self, tmp_path):
        doc = Document()
        doc.add_paragraph("body text")
        dm = DocxMap(doc)
        assert not hasattr(dm, "_metadata_settings")

        out_path = str(tmp_path / "out.docx")
        with mock.patch.object(DocxMap, "_rewrite_docx_zip") as rewrite:
            dm.save(out_path)
        rewrite.assert_not_called()
        assert os.path.exists(out_path)
